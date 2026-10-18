"""Write + read benchmark: xlspy vs xlsxwriter vs openpyxl.

Dataset: 50000 x 50 (2.5M cells) by default, mixed types
(str, int, float, bool, date, datetime, Decimal, big int).

Write:
  - xlspy (C_EXT) XLSB            - default, compiled C extension
  - xlspy (Python) XLSB           - pure-Python fallback, measured in a
                                    subprocess with XLSPY_DISABLE_C_EXT=1
  - xlspy XLSX
  - openpyxl XLSX (write_only=True, the streaming equivalent of
    xlsxwriter's constant_memory mode)
  - xlsxwriter XLSX (constant_memory=True)

Read (each library reads back its own output file):
  - xlspy (C_EXT) XLSB
  - xlspy (Python) XLSB           - subprocess with XLSPY_DISABLE_C_EXT=1
  - xlspy XLSX
  - openpyxl XLSX (read_only=True)

Usage:
    python examples/performance_test.py [--rows N] [--cols N] [--no-cleanup]
"""
import argparse
import os
import platform
import subprocess
import sys
import time

import xlsxwriter
try:
    from memory_profiler import profile
except ImportError:
    profile = None
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))
from xlspy import XlsbWriter
from xlspy import XlsxWriter
from xlspy import ExcelReader

import datetime
from decimal import Decimal

SHEET_NAME = "LargeSheet"


def generate_large_data(rows, cols):
    """Generates a large list of lists for testing."""
    print(f"Generating {rows}x{cols} dataset...")

    headers = [f"Header_{i}" for i in range(cols)]
    yield headers
    for i in range(1, rows):
        row_data = []
        for j in range(cols):
            col_type = j % 8
            if col_type == 0:
                row_data.append(f"String_{i}_{j}")
            elif col_type == 1:
                row_data.append(i * j)
            elif col_type == 2:
                row_data.append(i / (j + 1))
            elif col_type == 3:
                row_data.append(True if (i + j) % 2 == 0 else False)
            elif col_type == 4:
                row_data.append(datetime.date(2023, 1, 1) + datetime.timedelta(days=i))
            elif col_type == 5:
                row_data.append(datetime.datetime(2023, 1, 1, 12, 30, 0) + datetime.timedelta(seconds=i))
            elif col_type == 6:
                row_data.append(Decimal(f"{i}.{j}"))
            elif col_type == 7:
                row_data.append(2**60 + i)  # Large integer
        yield row_data


def _file_mb(filename):
    return os.path.getsize(filename) / (1024 * 1024)


# @profile
def run_my_xlsb_writer(data, filename="test_xlspy.xlsb"):
    """Tests the performance of XlsbWriter."""
    print("\n--- Testing (my) XlsbWriter ---")
    start_time = time.time()

    with XlsbWriter(filename) as writer:
        writer.add_sheet(SHEET_NAME)
        writer.write_sheet(data)


    end_time = time.time()
    elapsed = end_time - start_time

    filesize = _file_mb(filename)
    print(f"XlsbWriter finished in: {elapsed:.2f} seconds")
    print(f"File size: {filesize:.2f} MB")
    print("--------------------------")
    return elapsed, filesize


# @profile
def run_my_xlsx_writer(data, filename="test_xlspy.xlsx"):
    """Tests the performance of XlsxWriter."""
    print("\n--- Testing (my)  XlsxWriter ---")
    start_time = time.time()

    with XlsxWriter(filename) as writer:
        writer.add_sheet(SHEET_NAME)
        writer.write_sheet(data)


    end_time = time.time()
    elapsed = end_time - start_time

    filesize = _file_mb(filename)
    print(f"XlsxWriter finished in: {elapsed:.2f} seconds")
    print(f"File size: {filesize:.2f} MB")
    print("--------------------------")
    return elapsed, filesize


# @profile
def run_xlsx_writer(data, filename="test_xlsxwriter.xlsx"):
    """Tests the performance of xlsxwriter."""
    print("\n--- Testing xlsxwriter ---")
    start_time = time.time()

    # Use constant_memory mode for performance with large datasets
    workbook = xlsxwriter.Workbook(filename, {'constant_memory': True})
    worksheet = workbook.add_worksheet(SHEET_NAME)

    for row_idx, row_data in enumerate(data):
        worksheet.write_row(row_idx, 0, row_data)

    workbook.close()

    end_time = time.time()
    elapsed = end_time - start_time

    filesize = _file_mb(filename)
    print(f"xlsxwriter finished in: {elapsed:.2f} seconds")
    print(f"File size: {filesize:.2f} MB")
    print("--------------------------")
    return elapsed, filesize


def run_openpyxl_writer(data, filename="test_openpyxl.xlsx"):
    """Tests the performance of openpyxl in streaming (write-only) mode."""
    import openpyxl
    print("\n--- Testing openpyxl (write_only=True) ---")
    start_time = time.time()

    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet(SHEET_NAME)
    for row_data in data:
        worksheet.append(list(row_data))
    workbook.save(filename)

    end_time = time.time()
    elapsed = end_time - start_time

    filesize = _file_mb(filename)
    print(f"openpyxl finished in: {elapsed:.2f} seconds")
    print(f"File size: {filesize:.2f} MB")
    print("--------------------------")
    return elapsed, filesize


def run_read_xlspy(filename, sheet=SHEET_NAME):
    """Reads a file back with xlspy (ExcelReader), streaming rows."""
    print(f"\n--- Reading {filename} with xlspy ---")
    start_time = time.time()

    n_rows = 0
    with ExcelReader(filename) as reader:
        for _row in reader.get_rows(sheet):
            n_rows += 1

    end_time = time.time()
    elapsed = end_time - start_time
    print(f"xlspy read {n_rows} rows in: {elapsed:.2f} seconds")
    print("--------------------------")
    return elapsed, n_rows


def run_read_openpyxl(filename):
    """Reads a file back with openpyxl in read-only mode."""
    import openpyxl
    print(f"\n--- Reading {filename} with openpyxl (read_only=True) ---")
    start_time = time.time()

    n_rows = 0
    workbook = openpyxl.load_workbook(filename, read_only=True, data_only=True)
    worksheet = workbook[SHEET_NAME]
    for _row in worksheet.iter_rows(values_only=True):
        n_rows += 1
    workbook.close()

    end_time = time.time()
    elapsed = end_time - start_time
    print(f"openpyxl read {n_rows} rows in: {elapsed:.2f} seconds")
    print("--------------------------")
    return elapsed, n_rows


def _worker_write_xlsb_python(filename, rows, cols):
    """Worker entry point: times a pure-Python XLSB write.

    Must run in a subprocess with XLSPY_DISABLE_C_EXT=1 (the flag is
    evaluated at import time, so it cannot be toggled in-process).
    Prints `RESULT seconds=<t> size_mb=<s>` for the parent to parse.
    """
    elapsed, filesize = run_my_xlsb_writer(
        generate_large_data(rows, cols), filename
    )
    print(f"RESULT seconds={elapsed:.4f} size_mb={filesize:.4f}", flush=True)


def _worker_read_xlsb_python(filename):
    """Worker entry point: times a pure-Python XLSB read (subprocess)."""
    elapsed, n_rows = run_read_xlspy(filename)
    print(f"RESULT seconds={elapsed:.4f} rows={n_rows}", flush=True)


def _run_python_fallback_worker(worker, *args):
    """Runs this script as a subprocess with the C extension disabled."""
    env = dict(os.environ, XLSPY_DISABLE_C_EXT="1")
    cmd = [sys.executable, os.path.abspath(__file__), "--worker", worker,
           *args]
    proc = subprocess.run(cmd, capture_output=True, text=True, env=env)
    print(proc.stdout, end="")
    if proc.returncode != 0:
        print(proc.stderr, end="")
        raise RuntimeError(f"worker {worker} failed (exit {proc.returncode})")
    seconds = None
    extra = {}
    for line in proc.stdout.splitlines():
        if line.startswith("RESULT "):
            for token in line.split()[1:]:
                key, _, value = token.partition("=")
                if key == "seconds":
                    seconds = float(value)
                else:
                    extra[key] = float(value)
    if seconds is None:
        raise RuntimeError(f"worker {worker} printed no RESULT line")
    return seconds, extra


def _versions():
    import xlspy
    versions = {"python": platform.python_version()}
    try:
        versions["xlspy"] = xlspy.__version__
    except AttributeError:
        versions["xlspy"] = "?"
    for pkg in ("openpyxl", "xlsxwriter"):
        try:
            mod = __import__(pkg)
            versions[pkg] = getattr(mod, "__version__", "?")
        except ImportError:
            versions[pkg] = "missing"
    return versions


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="xlspy write/read benchmark")
    parser.add_argument("--rows", type=int, default=50000)
    parser.add_argument("--cols", type=int, default=50)
    parser.add_argument("--no-cleanup", action="store_true",
                        help="keep the generated benchmark files")
    parser.add_argument("--worker", choices=("write-xlsb-python", "read-xlsb-python"),
                        help=argparse.SUPPRESS)
    parser.add_argument("--out", default=None, help=argparse.SUPPRESS)
    args = parser.parse_args()

    if args.worker == "write-xlsb-python":
        _worker_write_xlsb_python(args.out, args.rows, args.cols)
        sys.exit(0)
    if args.worker == "read-xlsb-python":
        _worker_read_xlsb_python(args.out)
        sys.exit(0)

    # Define the size of the dataset
    NUM_ROWS = args.rows
    NUM_COLS = args.cols

    files_to_clean = []

    # ---- Write benchmarks ----
    f1 = "test_xlspy.xlsb"
    w_xlsb_c, s_xlsb = run_my_xlsb_writer(generate_large_data(NUM_ROWS, NUM_COLS), f1)
    files_to_clean.append(f1)

    print("\n--- Testing (my) XlsbWriter (pure Python fallback, subprocess) ---")
    w_xlsb_py, extra = _run_python_fallback_worker(
        "write-xlsb-python", "--rows", str(NUM_ROWS), "--cols", str(NUM_COLS),
        "--out", "test_xlspy_py.xlsb")
    s_xlsb_py = extra["size_mb"]
    print(f"XlsbWriter (Python) finished in: {w_xlsb_py:.2f} seconds")
    print(f"File size: {s_xlsb_py:.2f} MB")
    print("--------------------------")
    files_to_clean.append("test_xlspy_py.xlsb")

    f2 = "test_xlspy.xlsx"
    w_xlsx, s_xlsx = run_my_xlsx_writer(generate_large_data(NUM_ROWS, NUM_COLS), f2)
    files_to_clean.append(f2)

    f3 = "test_xlsxwriter.xlsx"
    w_xlsxwriter, s_xlsxwriter = run_xlsx_writer(generate_large_data(NUM_ROWS, NUM_COLS), f3)
    files_to_clean.append(f3)

    f4 = "test_openpyxl.xlsx"
    w_openpyxl, s_openpyxl = run_openpyxl_writer(generate_large_data(NUM_ROWS, NUM_COLS), f4)
    files_to_clean.append(f4)

    # ---- Read benchmarks (each library reads back its own output) ----
    r_xlsb_c, n_xlsb = run_read_xlspy(f1)
    print("\n--- Reading XLSB with xlspy (pure Python fallback, subprocess) ---")
    r_xlsb_py, extra = _run_python_fallback_worker("read-xlsb-python", "--out", f1)
    print(f"xlspy (Python) read {int(extra['rows'])} rows in: {r_xlsb_py:.2f} seconds")
    print("--------------------------")
    r_xlsx, n_xlsx = run_read_xlspy(f2)
    r_openpyxl, n_openpyxl = run_read_openpyxl(f4)

    # ---- Markdown-ready summary ----
    ver = _versions()
    print("\n================ SUMMARY ================")
    print(f"Dataset: {NUM_ROWS} x {NUM_COLS} "
          f"({NUM_ROWS * NUM_COLS / 1e6:.1f}M cells); "
          f"{platform.system()} {platform.machine()} "
          f"(Python {ver['python']}); "
          f"xlspy {ver['xlspy']}, openpyxl {ver['openpyxl']}, "
          f"xlsxwriter {ver['xlsxwriter']}")
    print("\n### Write\n")
    print("| Library | Format | Time | Size |")
    print("|---------|--------|------|------|")
    print(f"| **xlspy (C_EXT)** | XLSB | **{w_xlsb_c:.2f} s** | {s_xlsb:.2f} MB |")
    print(f"| xlspy (Python) | XLSB | {w_xlsb_py:.2f} s | {s_xlsb_py:.2f} MB |")
    print(f"| xlspy | XLSX | {w_xlsx:.2f} s | {s_xlsx:.2f} MB |")
    print(f"| openpyxl (write_only) | XLSX | {w_openpyxl:.2f} s | {s_openpyxl:.2f} MB |")
    print(f"| xlsxwriter (constant_memory) | XLSX | {w_xlsxwriter:.2f} s | {s_xlsxwriter:.2f} MB |")
    print("\n### Read\n")
    print("| Library | Format | Time | Notes |")
    print("|---------|--------|------|-------|")
    print(f"| **xlspy (C_EXT)** | XLSB | **{r_xlsb_c:.2f} s** | default, compiled C |")
    print(f"| xlspy | XLSX | {r_xlsx:.2f} s | uses expat XML parser (C) |")
    print(f"| xlspy (Python) | XLSB | {r_xlsb_py:.2f} s | pure Python fallback |")
    print(f"| openpyxl (read_only) | XLSX | {r_openpyxl:.2f} s | read-only mode |")
    print("=========================================")

    # Cleanup
    if not args.no_cleanup:
        for f in files_to_clean:
            try:
                os.remove(f)
            except OSError:
                pass
        print("Temporary benchmark files removed.")

    print("\nPerformance test complete.")
    print("To see detailed memory usage, run this script with memory-profiler:")
    print("mprof run performance_test.py")
    print("Then, to see the plot:")
    print("mprof plot")
