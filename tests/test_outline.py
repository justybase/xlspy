import struct
import zipfile

import openpyxl
import pytest

from xlspy import ExcelReader, Group, XlsbWriter, XlsxWriter
from xlspy.biff_utils import iter_records
from xlspy.outline import MAX_LEVEL, build_outline


ROWS = [[f"r{i}c{j}" for j in range(10)] for i in range(12)]
ROW_GROUPS = [Group(2, 7), Group(2, 3), Group(5, 6, collapsed=True)]
COL_GROUPS = [Group("D", "F"), Group("H", "J")]


def _sheet_payloads(path, record_id, part="xl/worksheets/sheet1.bin"):
    with zipfile.ZipFile(path) as archive:
        blob = archive.read(part)
    records = list(iter_records(blob))
    assert records[-1].data_end == len(blob), "sheet records must tile the buffer"
    return [blob[record.data_start:record.data_end] for record in records if record.record_id == record_id]


def _row_flags(path):
    """Map zero-based row index to the BrtRowHdr flags byte (payload offset 11)."""
    return {
        struct.unpack_from("<I", payload, 0)[0]: payload[11]
        for payload in _sheet_payloads(path, 0x0000)
    }


def test_build_outline_returns_none_without_groups():
    assert build_outline(None, "row") is None
    assert build_outline([], "column") is None


def test_outline_levels_follow_nesting_depth():
    outline = build_outline(ROW_GROUPS, "row")

    assert outline.max_level == 2
    assert [outline.state_at(i) for i in range(9)] == [
        (0, False), (0, False),
        (2, False), (2, False), (1, False),
        (2, True), (2, True), (1, False), (0, False),
    ]
    assert outline.collapsed_at(7)
    assert not outline.collapsed_at(6)
    assert outline.indices() == {2, 3, 4, 5, 6, 7}
    assert outline.hidden_indices() == {5, 6}


def test_collapsed_outer_group_hides_everything_inside_it():
    outline = build_outline([Group(0, 4, collapsed=True), Group(1, 2)], "row")

    assert outline.hidden_indices() == {0, 1, 2, 3, 4}
    assert outline.collapsed_at(5)


def test_tuples_and_letters_are_accepted():
    rows = build_outline([(0, 1), (3, 4, True)], "row")
    assert rows.state_at(1) == (1, False)
    assert rows.collapsed_at(5)

    columns = build_outline([Group("D", "F"), ("aa", "ab")], "column")
    assert columns.state_at(3) == (1, False)
    assert columns.state_at(26) == (1, False)


@pytest.mark.parametrize(
    ("groups", "kind", "error", "message"),
    [
        ([Group(1, 0)], "row", ValueError, "ends before it starts"),
        ([Group(0, 5), Group(3, 8)], "row", ValueError, "overlap without nesting"),
        ([Group(0, 3), Group(0, 3)], "row", ValueError, "duplicate row group"),
        ([Group("A", 3)], "row", TypeError, "row group bounds must be 0-based ints"),
        ([Group(True, 3)], "row", TypeError, "row group bounds must be 0-based ints"),
        ([Group(0, 1048576)], "row", ValueError, "out of range"),
        ([Group("XFE", "XFE")], "column", ValueError, "out of range"),
        ([Group("A1", "B")], "column", TypeError, "column group bounds must be"),
        ([Group(0, 1, "yes")], "column", TypeError, "collapsed flags must be bool"),
        ([3], "row", TypeError, "entries must be Group objects"),
        ("abc", "row", TypeError, "not a string"),
        (5, "row", TypeError, "must be a collection"),
        ([Group(i, 20 - i) for i in range(MAX_LEVEL + 1)], "row", ValueError, "nest deeper"),
    ],
)
def test_invalid_groups_are_rejected(groups, kind, error, message):
    with pytest.raises(error, match=message):
        build_outline(groups, kind)


def test_adjacent_groups_share_level_one():
    outline = build_outline([Group(0, 1), Group(2, 3)], "row")
    assert outline.max_level == 1
    assert outline.indices() == {0, 1, 2, 3}


def test_xlsx_writes_outline_attributes_and_reads_back(tmp_path):
    path = tmp_path / "grouped.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, row_groups=ROW_GROUPS, column_groups=COL_GROUPS)

    sheet = openpyxl.load_workbook(path)["S"]
    levels = {r: sheet.row_dimensions[r].outline_level for r in range(1, 13)}
    assert {r: level for r, level in levels.items() if level} == {3: 2, 4: 2, 5: 1, 6: 2, 7: 2, 8: 1}
    assert [r for r in range(1, 13) if sheet.row_dimensions[r].hidden] == [6, 7]
    assert [r for r in range(1, 13) if sheet.row_dimensions[r].collapsed] == [8]
    assert sheet.sheet_format.outlineLevelRow == 2
    assert sheet.sheet_format.outlineLevelCol == 1
    assert {c: sheet.column_dimensions[c].outline_level for c in "DEFHIJ"} == {
        "D": 1, "E": 1, "F": 1, "H": 1, "I": 1, "J": 1,
    }


def test_xlsx_without_groups_has_no_outline_markup(tmp_path):
    path = tmp_path / "plain.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, hidden_rows=[2])

    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml").decode()
    assert "outline" not in xml
    assert "collapsed" not in xml
    assert "sheetFormatPr" not in xml


def test_xlsx_column_summary_collapse_is_written(tmp_path):
    path = tmp_path / "collapsed_cols.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, column_groups=[Group("D", "F", collapsed=True)])

    sheet = openpyxl.load_workbook(path)["S"]
    assert [c for c in "DEF" if sheet.column_dimensions[c].hidden] == ["D", "E", "F"]
    assert sheet.column_dimensions["G"].collapsed


def test_xlsx_streams_generator_rows(tmp_path):
    path = tmp_path / "stream.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet((row for row in ROWS), row_groups=ROW_GROUPS)

    sheet = openpyxl.load_workbook(path)["S"]
    assert sheet.row_dimensions[3].outline_level == 2
    assert sheet.row_dimensions[7].hidden is True


def test_xlsb_row_headers_carry_level_hidden_and_collapsed_bits(tmp_path):
    path = tmp_path / "grouped.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, row_groups=ROW_GROUPS, column_groups=COL_GROUPS)

    flags = _row_flags(path)
    assert flags[2] == 0x02          # outline level 2
    assert flags[3] == 0x02
    assert flags[4] == 0x01          # outline level 1, summary of 2..3
    assert flags[5] == 0x12          # level 2, hidden
    assert flags[6] == 0x12
    assert flags[7] == 0x09          # level 1, collapsed summary row
    assert flags[0] == 0 and flags[1] == 0 and flags[8] == 0


def test_xlsb_column_infos_carry_level_and_collapsed_bits(tmp_path):
    path = tmp_path / "grouped_cols.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, column_groups=[Group("D", "F", collapsed=True)])

    columns = {
        first: flags
        for first, last, _, _, flags in (
            struct.unpack("<IIIIH", payload) for payload in _sheet_payloads(path, 0x003C)
        )
    }
    assert columns[2] == 0x0002          # column C: ungrouped
    assert columns[3] == 0x0103          # column D: level 1, hidden
    assert columns[5] == 0x0103          # column F: level 1, hidden
    assert columns[6] == 0x1002          # column G: summary, collapsed


def test_xlsb_nested_column_levels(tmp_path):
    path = tmp_path / "nested_cols.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, column_groups=[Group(2, 5), Group(3, 4)])

    levels = {
        first: (flags >> 8) & 7
        for first, _, _, _, flags in (
            struct.unpack("<IIIIH", payload) for payload in _sheet_payloads(path, 0x003C)
        )
    }
    assert [levels[c] for c in range(2, 6)] == [1, 2, 2, 1]


def test_xlsb_column_groups_past_the_data_get_column_infos(tmp_path):
    path = tmp_path / "beyond.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet([["a", "b"]], column_groups=[Group("E", "F")])

    columns = {
        first: (last, width, flags)
        for first, last, width, _, flags in (
            struct.unpack("<IIIIH", payload) for payload in _sheet_payloads(path, 0x003C)
        )
    }
    assert columns[4] == (4, 8 * 256, 0x0102)
    assert columns[5] == (5, 8 * 256, 0x0102)


def test_xlsb_format_info_holds_max_outline_levels(tmp_path):
    grouped = tmp_path / "grouped.xlsb"
    with XlsbWriter(grouped) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, row_groups=ROW_GROUPS, column_groups=COL_GROUPS)
    fmt = _sheet_payloads(grouped, 0x01E5)[0]
    assert (fmt[10], fmt[11]) == (2, 1)

    plain = tmp_path / "plain.xlsb"
    with XlsbWriter(plain) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS)
    fmt = _sheet_payloads(plain, 0x01E5)[0]
    assert (fmt[10], fmt[11]) == (0, 0)


def test_xlsb_plain_sheet_records_are_unchanged_by_the_feature(tmp_path):
    path = tmp_path / "plain.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, hidden_rows=[2])

    flags = _row_flags(path)
    assert flags[2] == 0x10          # only the existing hidden bit
    assert all(flags[i] == 0 for i in range(12) if i != 2)


@pytest.mark.parametrize("writer_cls, ext", [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")])
def test_grouped_sheets_read_back_their_cells(tmp_path, writer_cls, ext):
    path = tmp_path / f"grouped.{ext}"
    with writer_cls(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet(ROWS, row_groups=ROW_GROUPS, column_groups=COL_GROUPS)

    with ExcelReader(str(path)) as reader:
        assert list(reader.get_rows("S")) == ROWS


def test_to_xlsx_passes_groups_through(tmp_path):
    pd = pytest.importorskip("pandas")
    from xlspy import to_xlsx

    df = pd.DataFrame(ROWS[1:], columns=ROWS[0])
    path = tmp_path / "frame.xlsx"
    to_xlsx(df, path, row_groups=[Group(0, 2, collapsed=True)])

    sheet = openpyxl.load_workbook(path).active
    assert sheet.row_dimensions[2].hidden is True
    assert sheet.row_dimensions[4].collapsed is True
