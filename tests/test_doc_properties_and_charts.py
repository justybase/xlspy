"""Document properties, stacked/area/doughnut charts and the XLSX spool size."""

import datetime
import zipfile
from xml.etree import ElementTree

import openpyxl
import pytest

from xlspy import Chart, ChartSeries, XlsbWriter, XlsxWriter, load_workbook
from xlspy.xlsx_writer import DEFAULT_SPOOL_BYTES

ROWS = [["Month", "North", "South"], ["Jan", 10, 20], ["Feb", 15, 25], ["Mar", 12, 30]]
SERIES = [
    ChartSeries(values="B2:B4", categories="A2:A4", title_cell="B1"),
    ChartSeries(values="C2:C4", categories="A2:A4", title_cell="C1"),
]
WRITERS = [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")]


def _write(path, writer_cls, **props):
    with writer_cls(path) as writer:
        if props:
            writer.set_properties(**props)
        writer.add_sheet("Data")
        writer.write_sheet(ROWS)


@pytest.mark.parametrize("writer_cls,ext", WRITERS)
def test_properties_parts(tmp_path, writer_cls, ext):
    path = tmp_path / f"props.{ext}"
    _write(
        path, writer_cls, title="Q3 & <report>", author="Zoë Ångström", company="ACME",
        manager="Boss", keywords="sales, q3", comments="notes", category="Finance",
        status="Draft", subject="Sales",
        created=datetime.datetime(2024, 5, 1, 10, 0, tzinfo=datetime.timezone.utc),
        custom={"Project": "X", "Count": 7, "Big": 2**40, "Ratio": 0.5, "Approved": True,
                "Due": datetime.date(2024, 6, 1)},
    )
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        core = archive.read("docProps/core.xml").decode()
        app = archive.read("docProps/app.xml").decode()
        custom = archive.read("docProps/custom.xml").decode()
        types = archive.read("[Content_Types].xml").decode()
        rels = archive.read("_rels/.rels").decode()
    assert {"docProps/core.xml", "docProps/app.xml", "docProps/custom.xml"} <= set(names)
    assert "<dc:title>Q3 &amp; &lt;report&gt;</dc:title>" in core
    assert "<dc:creator>Zoë Ångström</dc:creator>" in core
    assert '<dcterms:created xsi:type="dcterms:W3CDTF">2024-05-01T10:00:00Z' in core
    assert "<Company>ACME</Company>" in app and "<Manager>Boss</Manager>" in app
    for fragment in ('name="Project"><vt:lpwstr>X', "<vt:i4>7</vt:i4>", "<vt:r8>1099511627776.0",
                     "<vt:r8>0.5</vt:r8>", "<vt:bool>true", "<vt:filetime>2024-06-01T00:00:00Z"):
        assert fragment in custom
    for part in ("core.xml", "app.xml", "custom.xml"):
        assert f"/docProps/{part}" in types and f"docProps/{part}" in rels


def test_properties_read_back(tmp_path):
    path = tmp_path / "props.xlsx"
    _write(path, XlsxWriter, title="T", author="A", keywords="k", category="c",
           created=datetime.datetime(2024, 1, 2, 3, 4, 5))
    props = openpyxl.load_workbook(path).properties
    assert (props.title, props.creator, props.keywords, props.category) == ("T", "A", "k", "c")
    assert props.created == datetime.datetime(2024, 1, 2, 3, 4, 5)
    own = load_workbook(path)
    assert own.properties.title == "T"
    own.close()


@pytest.mark.parametrize("writer_cls,ext", WRITERS)
def test_no_properties_parts_unless_asked(tmp_path, writer_cls, ext):
    path = tmp_path / f"plain.{ext}"
    _write(path, writer_cls)
    with zipfile.ZipFile(path) as archive:
        assert not any(name.startswith("docProps/") for name in archive.namelist())
        assert "docProps" not in archive.read("_rels/.rels").decode()


def test_properties_validation(tmp_path):
    with XlsxWriter(tmp_path / "bad.xlsx") as writer:
        with pytest.raises(TypeError, match="title must be a str"):
            writer.set_properties(title=3)
        with pytest.raises(ValueError, match="control character"):
            writer.set_properties(title="a\x01b")
        with pytest.raises(TypeError, match="created"):
            writer.set_properties(created="2024-01-01")
        with pytest.raises(TypeError, match="custom property 'x'"):
            writer.set_properties(custom={"x": [1]})
        with pytest.raises(ValueError, match="duplicate"):
            writer.set_properties(custom={"Name": 1, "name": 2})
        with pytest.raises(ValueError, match="finite"):
            writer.set_properties(custom={"x": float("nan")})
        writer.add_sheet("S")
        writer.write_sheet([[1]])


def test_workbook_model_saves_its_properties(tmp_path):
    from xlspy import Workbook

    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.properties.title = "Model title"
    workbook.properties.creator = "Model author"
    path = tmp_path / "model.xlsx"
    workbook.save(path)
    props = openpyxl.load_workbook(path).properties
    assert (props.title, props.creator) == ("Model title", "Model author")


@pytest.mark.parametrize("chart_type,grouping,tag,expected", [
    ("bar", "stacked", "barChart", 'grouping val="stacked"'),
    ("barh", "percent_stacked", "barChart", 'grouping val="percentStacked"'),
    ("line", "stacked", "lineChart", 'grouping val="stacked"'),
    ("area", None, "areaChart", 'grouping val="standard"'),
    ("area", "percent_stacked", "areaChart", 'grouping val="percentStacked"'),
    ("doughnut", None, "doughnutChart", 'holeSize val="50"'),
])
def test_new_chart_types_and_grouping(tmp_path, chart_type, grouping, tag, expected):
    chart = Chart(chart_type, SERIES[:1] if chart_type == "doughnut" else SERIES,
                  title="Sales", grouping=grouping)
    path = tmp_path / "chart.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS, charts=[chart])
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("xl/charts/chart1.xml").decode()
    assert f"<{tag}>" in xml and expected in xml
    if chart_type in ("bar", "barh"):
        assert '<overlap val="100"/>' in xml
    ElementTree.fromstring(xml)  # well-formed; Excel itself checks the rest (COM tests)


def test_chart_grouping_validation():
    with pytest.raises(ValueError, match="invalid chart grouping"):
        Chart("bar", SERIES, grouping="3d")
    with pytest.raises(ValueError, match="only valid for"):
        Chart("pie", SERIES[:1], grouping="stacked")
    with pytest.raises(ValueError, match="no axes"):
        from xlspy import ChartAxis
        Chart("doughnut", SERIES[:1], y_axis=ChartAxis(minimum=0))
    with pytest.raises(ValueError, match="series color"):
        Chart("doughnut", [ChartSeries(values="B2:B4", color="#FF0000")])
    with pytest.raises(ValueError, match="marker"):
        from xlspy import ChartMarker
        Chart("area", [ChartSeries(values="B2:B4", marker=ChartMarker())])


def test_default_spool_is_small():
    # Bigger spools only cost memory (measured: 8 MiB vs 1 MiB, same speed).
    assert DEFAULT_SPOOL_BYTES == 1 << 20
