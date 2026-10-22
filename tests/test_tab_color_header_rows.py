"""Sheet tab colours and headers of several rows (filter, header style)."""

import re
import struct
import zipfile

import openpyxl
import pytest

from xlspy import Color, PrintSettings, Style, Workbook, XlsbWriter, XlsxWriter
from xlspy.biff_utils import iter_records

BRT_WS_PROP = 0x0093
BRT_BEGIN_AFILTER = 0x00A1
HEADER = Style(fill="#1F4E78", font_color="#FFFFFF")
ROWS = [
    ["Region", None, "Sales", None],
    ["Name", "Code", "Q1", "Q2"],
    ["North", "N", 1, 2],
    ["South", "S", 3, 4],
]


def _xlsx_sheet(path, index=1):
    with zipfile.ZipFile(path) as archive:
        return archive.read(f"xl/worksheets/sheet{index}.xml").decode()


def _xlsb_records(path):
    with zipfile.ZipFile(path) as archive:
        data = archive.read("xl/worksheets/sheet1.bin")
    return {record.record_id: data[record.data_start:record.data_end]
            for record in iter_records(data)}


def test_xlsx_tab_colors(tmp_path):
    path = tmp_path / "tabs.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("Rgb")
        writer.write_sheet(ROWS, tab_color="#C00000",
                           print_settings=PrintSettings(fit_to_width=1))
        writer.add_sheet("Theme")
        writer.write_sheet(ROWS, tab_color=Color(theme=5, tint=0.4))
        writer.add_sheet("Plain")
        writer.write_sheet(ROWS)
    # tabColor comes before pageSetUpPr inside one sheetPr.
    assert '<sheetPr><tabColor rgb="FFC00000"/><pageSetUpPr fitToPage="1"/></sheetPr>' in _xlsx_sheet(path, 1)
    assert '<sheetPr><tabColor theme="5" tint="0.4"/></sheetPr>' in _xlsx_sheet(path, 2)
    assert "<sheetPr" not in _xlsx_sheet(path, 3)
    workbook = openpyxl.load_workbook(path)
    assert workbook["Rgb"].sheet_properties.tabColor.rgb == "FFC00000"
    workbook.close()


def test_xlsb_tab_color(tmp_path):
    path = tmp_path / "tab.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS, tab_color="C00000")
    props = _xlsb_records(path)[BRT_WS_PROP]
    # brtcolorTab: RGB colour type with fValidRGB, then red, green, blue, alpha.
    assert props[3:11] == bytes((0x05, 0, 0, 0, 0xC0, 0, 0, 0xFF))


def test_tab_color_validation(tmp_path):
    with XlsxWriter(tmp_path / "bad.xlsx") as writer:
        writer.add_sheet("Data")
        with pytest.raises(ValueError):
            writer.write_sheet(ROWS, tab_color="not a colour")
        with pytest.raises(TypeError, match="tab_color"):
            writer.write_sheet(ROWS, tab_color=12)
        writer.write_sheet(ROWS)


def test_xlsx_filter_and_style_on_a_two_row_header(tmp_path):
    path = tmp_path / "header.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS, header_rows=2, header_style=HEADER, merges=["A1:B1", "C1:D1"],
                           styles={"C2:D4": Style(number_format="0.00")})
    xml = _xlsx_sheet(path)
    assert '<autoFilter ref="A2:D4"/>' in xml
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Data"]
    assert sheet.auto_filter.ref == "A2:D4"
    for ref in ("A1", "C1", "A2", "D2"):  # B1 and D1 are merged away
        assert sheet[ref].fill.fgColor.rgb == "FF1F4E78", ref
        assert sheet[ref].font.bold and sheet[ref].font.color.rgb == "FFFFFFFF", ref
    # A range on the header overlays the header style, as with one header row.
    assert sheet["C2"].number_format == "0.00" and sheet["C2"].fill.fgColor.rgb == "FF1F4E78"
    assert sheet["A3"].fill.fill_type is None and not sheet["A3"].font.bold
    workbook.close()
    defined = re.search(r"_xlnm\._FilterDatabase[^>]*>([^<]+)<", zipfile.ZipFile(path).read("xl/workbook.xml").decode())
    assert defined.group(1) == "'Data'!$A$2:$D$4"


def test_xlsb_filter_on_a_two_row_header(tmp_path):
    path = tmp_path / "header.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS, header_rows=2, header_style=HEADER)
    first_row, last_row, first_col, last_col = struct.unpack("<IIII", _xlsb_records(path)[BRT_BEGIN_AFILTER])
    assert (first_row, last_row, first_col, last_col) == (1, 3, 0, 3)


@pytest.mark.parametrize("writer_cls,ext", [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")])
def test_one_header_row_is_unchanged(tmp_path, writer_cls, ext):
    default, explicit = tmp_path / f"a.{ext}", tmp_path / f"b.{ext}"
    for path, options in ((default, {}), (explicit, {"header_rows": 1})):
        with writer_cls(path) as writer:
            writer.add_sheet("Data")
            writer.write_sheet(ROWS, header_style=HEADER, **options)
    with zipfile.ZipFile(default) as a, zipfile.ZipFile(explicit) as b:
        for name in a.namelist():
            assert a.read(name) == b.read(name), name


def test_header_taller_than_the_data_gets_no_filter(tmp_path):
    path = tmp_path / "short.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS[:2], header_rows=2)
    assert "<autoFilter" not in _xlsx_sheet(path)


def test_header_rows_validation(tmp_path):
    with XlsxWriter(tmp_path / "bad.xlsx") as writer:
        writer.add_sheet("Data")
        with pytest.raises(ValueError, match="at least 1"):
            writer.write_sheet(ROWS, header_rows=0)
        with pytest.raises(TypeError, match="header_rows"):
            writer.write_sheet(ROWS, header_rows=True)
        writer.write_sheet(ROWS)


def test_workbook_model_saves_tab_color(tmp_path):
    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.active.sheet_properties.tabColor = "00B050"
    path = tmp_path / "model.xlsx"
    workbook.save(path)
    loaded = openpyxl.load_workbook(path)
    assert loaded.active.sheet_properties.tabColor.rgb == "FF00B050"
    loaded.close()
