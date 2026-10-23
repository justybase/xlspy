"""openpyxl-style scripts on the workbook model: objects, printing, zoom,
row/column moves, rich text and the extra conditional formatting rules."""

import datetime
import io
import re
import struct
import zipfile
import zlib

import openpyxl
import pytest

import xlspy

from xlspy import (
    AboveAverageRule, BlanksRule, CellIsRule, CellRichText, Chart, ChartSeries, ExcelReader,
    DataValidation, DifferentialStyle, DuplicateRule, ExcelTable, Font, Image,
    InlineFont, TextBlock, TextRule, TimePeriodRule, TopBottomRule, Workbook,
    Style, XlsbWriter, XlsxWriter, load_workbook,
)
from xlspy.workbook._shift import translate_formula

BOLD = DifferentialStyle(font=Font(bold=True))


def _png(width, height):
    raw = b"".join(b"\x00" + b"\x10\x80\x10" * width for _ in range(height))

    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))

    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _sheet_xml(path, index=1):
    with zipfile.ZipFile(path) as archive:
        return archive.read(f"xl/worksheets/sheet{index}.xml").decode()


def _report_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.append(["Region", "Sales"])
    for row in (["North", 10], ["South", 30], ["East", 20]):
        sheet.append(row)
    return workbook, sheet


# ----------------------------------------------------------------- objects


def test_model_objects_are_saved(tmp_path):
    workbook, sheet = _report_workbook()
    sheet.add_chart(Chart("bar", [ChartSeries("B2:B4", categories="A2:A4")]), "D2")
    sheet.add_image(Image(io.BytesIO(_png(8, 4))), "D20")
    sheet.add_table(ExcelTable(display_name="Sales", ref="A1:B4"))
    sheet.add_data_validation(DataValidation(type="whole", sqref="B2:B4", operator="between",
                                             formula1="0", formula2="100"))
    sheet.conditional_formatting.add("B2:B4", CellIsRule("A1", "greaterThan", ["15"], dxf=BOLD))
    path = tmp_path / "objects.xlsx"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    assert {"xl/charts/chart1.xml", "xl/tables/table1.xml", "xl/media/image1.png"} <= set(names)
    loaded = openpyxl.load_workbook(path)["Report"]
    assert list(loaded.tables) == ["Sales"]
    assert str(loaded.data_validations.dataValidation[0].sqref) == "B2:B4"
    assert [str(cf.sqref) for cf in loaded.conditional_formatting] == ["B2:B4"]
    assert re.search(r'<drawing r:id="rId\d+"/>', _sheet_xml(path))


def test_model_objects_reject_openpyxl_classes_and_save_to_xlsb(tmp_path):
    from openpyxl.chart import BarChart
    from xlspy.biff_utils import iter_records

    workbook, sheet = _report_workbook()
    with pytest.raises(TypeError, match="use an xlspy Chart"):
        sheet.add_chart(BarChart())
    with pytest.raises(TypeError, match="DataValidation"):
        sheet.add_data_validation("A1")
    sheet.add_data_validation(DataValidation(type="list", sqref="A2:A4", formula1='"a,b"'))
    workbook.save(tmp_path / "x.xlsb")
    data = zipfile.ZipFile(tmp_path / "x.xlsb").read("xl/worksheets/sheet1.bin")
    assert [record.record_id for record in iter_records(data)].count(0x40) == 1  # BrtDVal


def test_loaded_workbook_saves_new_objects(tmp_path):
    workbook, _ = _report_workbook()
    workbook.save(tmp_path / "base.xlsx")
    loaded = load_workbook(tmp_path / "base.xlsx")
    loaded.active.add_table(ExcelTable(display_name="Sales", ref="A1:B4"))
    loaded.save(tmp_path / "out.xlsx")
    loaded.close()
    assert list(openpyxl.load_workbook(tmp_path / "out.xlsx").active.tables) == ["Sales"]


# ------------------------------------------------------------ print, zoom


def test_model_print_setup_and_zoom(tmp_path):
    workbook, sheet = _report_workbook()
    sheet.print_title_rows = "1:1"
    sheet.print_area = ["A1:B4"]
    sheet.page_setup.orientation = sheet.ORIENTATION_LANDSCAPE
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.4
    sheet.oddHeader.center.text = "Report"
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.sheet_view.zoomScale = 85
    path = tmp_path / "print.xlsx"
    workbook.save(path)
    loaded = openpyxl.load_workbook(path)["Report"]
    assert loaded.print_title_rows == "$1:$1"
    assert loaded.print_area == "'Report'!$A$1:$B$4"
    assert loaded.page_setup.orientation == "landscape" and int(loaded.page_setup.paperSize) == 9
    assert loaded.sheet_properties.pageSetUpPr.fitToPage
    assert (loaded.page_setup.fitToWidth, loaded.page_setup.fitToHeight) == (1, 0)
    assert loaded.print_options.horizontalCentered and loaded.page_margins.left == 0.4
    assert loaded.oddHeader.center.text == "Report"
    assert loaded.oddFooter.right.text == "Page &P of &N"
    assert loaded.sheet_view.zoomScale == 85


def test_untouched_print_setup_writes_nothing(tmp_path):
    workbook, _ = _report_workbook()
    workbook.save(tmp_path / "plain.xlsx")
    xml = _sheet_xml(tmp_path / "plain.xlsx")
    assert "<pageSetup" not in xml and "zoomScale" not in xml


@pytest.mark.parametrize("writer_cls,ext", [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")])
def test_writer_zoom(tmp_path, writer_cls, ext):
    path = tmp_path / f"zoom.{ext}"
    with writer_cls(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([[1]], zoom=150)
        with pytest.raises(ValueError, match="between 10 and 400"):
            writer.write_sheet([[1]], zoom=5)
    if ext == "xlsx":
        assert 'zoomScale="150" zoomScaleNormal="150"' in _sheet_xml(path)
    else:
        with zipfile.ZipFile(path) as archive:
            data = archive.read("xl/worksheets/sheet1.bin")
        view = data.index(bytes((0x89, 0x01, 0x1E))) + 3
        assert struct.unpack_from("<HH", data, view + 18) == (150, 150)


# ------------------------------------------------- insert / delete / move


def _values(sheet):
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(
        min_row=1, max_row=sheet.max_row, max_col=sheet.max_column)]
    while rows and all(value is None for value in rows[-1]):
        rows.pop()  # openpyxl counts cells emptied by a move
    return rows


def _shift_script(module):
    workbook = module.Workbook()
    sheet = workbook.active
    for r in range(1, 8):
        sheet.append([f"r{r}c{c}" for c in range(1, 6)])
    sheet["F1"] = "=SUM(A1:B2)+$C$3"
    sheet.insert_rows(2, 2)
    sheet.delete_rows(6)
    sheet.insert_cols(2)
    sheet.delete_cols(4, 2)
    sheet.move_range("A1:C3", rows=8, cols=1, translate=True)
    sheet.move_range("D1:D2", rows=1, cols=2, translate=True)
    return _values(sheet)


def test_moves_match_openpyxl():
    assert _shift_script(xlspy) == _shift_script(openpyxl)


def test_insert_rows_moves_styles_comments_and_heights():
    from xlspy import Comment

    workbook = Workbook()
    sheet = workbook.active
    sheet["A2"] = "x"
    sheet["A2"].font = Font(bold=True)
    sheet["A2"].comment = Comment("note", "me")
    sheet.row_dimensions[2].height = 30
    sheet.column_dimensions["B"].width = 20
    sheet.insert_rows(1, 3)
    sheet.insert_cols(1)
    assert sheet["B5"].value == "x" and sheet["B5"].font.bold and sheet["B5"].comment.text == "note"
    assert sheet.row_dimensions[5].height == 30 and sheet.row_dimensions.get(2) is None
    assert sheet.column_dimensions["C"].width == 20
    assert (sheet.max_row, sheet.max_column) == (5, 2)
    sheet.delete_rows(1, 4)
    assert sheet["B1"].value == "x" and sheet.max_row == 1


def test_translate_formula():
    assert translate_formula('=A1+$B$2+B$3+$C4+Sheet2!D5+"A1"', 2, 1) == '=B3+$B$2+C$3+$C6+Sheet2!E7+"A1"'
    assert translate_formula("=SUM(A1:B2)+LOG10(A1)", 1, 0) == "=SUM(A2:B3)+LOG10(A2)"
    assert translate_formula("=A1", -1, 0) == "=#REF!"
    assert translate_formula("='My A1'!A1", 0, 1) == "='My A1'!B1"


def test_shifts_in_a_loaded_workbook(tmp_path):
    source = openpyxl.Workbook()
    sheet = source.active
    sheet.title = "D"
    for r in range(1, 6):
        sheet.append([f"r{r}", r, r * 10])
    sheet["A2"].font = openpyxl.styles.Font(bold=True)
    sheet.row_dimensions[3].height = 30
    sheet.column_dimensions["C"].width = 25
    source.save(tmp_path / "src.xlsx")
    workbook = load_workbook(tmp_path / "src.xlsx")
    workbook["D"].insert_rows(2)
    workbook["D"].delete_cols(2)
    workbook["D"]["A2"] = "new"
    workbook.save(tmp_path / "out.xlsx")
    workbook.close()
    result = openpyxl.load_workbook(tmp_path / "out.xlsx")["D"]
    assert [[c.value for c in row] for row in result.iter_rows(max_col=2)] == [
        ["r1", 10], ["new", None], ["r2", 20], ["r3", 30], ["r4", 40], ["r5", 50],
    ]
    assert result["A3"].font.b and not result["A2"].font.b
    assert result.row_dimensions[4].height == 30 and result.row_dimensions[3].height is None
    assert result.column_dimensions["B"].width == 25


def test_shift_validation():
    sheet = Workbook().active
    with pytest.raises(ValueError, match="amount"):
        sheet.insert_rows(1, 0)
    with pytest.raises(TypeError):
        sheet.insert_cols("B")
    sheet["A1048576"] = 1
    with pytest.raises(ValueError, match="last row"):
        sheet.insert_rows(1)
    assert sheet["A1048576"].value == 1  # nothing moved


# ------------------------------------------------------------- rich text

RICH = CellRichText("Total: ", TextBlock(InlineFont(b=True, color="C00000", sz=14), "1 234"), " EUR")


@pytest.mark.parametrize("shared", [True, False])
def test_rich_text_cells(tmp_path, shared):
    path = tmp_path / "rich.xlsx"
    with XlsxWriter(path, useSharedStrings=shared) as writer:
        writer.add_sheet("S")
        writer.write_sheet([["plain", RICH], [RICH, "Total: 1 234 EUR"]])
    sheet = openpyxl.load_workbook(path, rich_text=True)["S"]
    value = sheet["B1"].value
    assert str(value) == "Total: 1 234 EUR" and sheet["A2"].value == value
    block = value[1]
    assert block.text == "1 234" and block.font.b and block.font.color.rgb == "FFC00000"
    assert block.font.sz == 14
    assert sheet["B2"].value == "Total: 1 234 EUR"  # plain text stays its own entry


def test_rich_text_in_the_model_and_loaded_files(tmp_path):
    workbook = Workbook()
    workbook.active["A1"] = RICH
    workbook.save(tmp_path / "model.xlsx")
    loaded = load_workbook(tmp_path / "model.xlsx", rich_text=True)
    assert loaded.active["A1"].value == RICH
    plain = load_workbook(tmp_path / "model.xlsx")
    assert plain.active["A1"].value == "Total: 1 234 EUR"
    plain.close()
    loaded.active["B2"] = RICH
    loaded.save(tmp_path / "edited.xlsx")
    loaded.close()
    edited = openpyxl.load_workbook(tmp_path / "edited.xlsx", rich_text=True).active
    assert edited["B2"].value[1].font.b


@pytest.mark.parametrize("shared", [True, False])
def test_rich_text_escapes_are_written_once(tmp_path, shared):
    """Runs already hold their _xHHHH_ escapes; the shared-string table must
    not escape them again."""
    rich = CellRichText("a\rb_x0041_", TextBlock(InlineFont(b=True), "c\x01"))
    path = tmp_path / "escaped.xlsx"
    with XlsxWriter(path, useSharedStrings=shared) as writer:
        writer.add_sheet("S")
        writer.write_sheet([["head"], [rich], ["d\re"]])
    loaded = load_workbook(path, rich_text=True)
    assert str(loaded["S"]["A2"].value) == "a\rb_x0041_c\x01"
    assert loaded["S"]["A3"].value == "d\re"
    loaded.close()


def test_rich_text_run_with_underline_none_loads(tmp_path):
    """Fonts other producers write (u val="none", an out-of-range size) must
    not make loading fail."""
    source = tmp_path / "source.xlsx"
    with XlsxWriter(source) as writer:
        writer.add_sheet("S")
        writer.write_sheet([[RICH]])
    path = tmp_path / "underline_none.xlsx"
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(path, "w") as archive:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/sharedStrings.xml":
                data = data.replace(b"<b/>", b'<b/><u val="none"/>')
                data = data.replace(b'<sz val="14"/>', b'<sz val="0.5"/>')
            archive.writestr(name, data)
    loaded = load_workbook(path, rich_text=True)
    block = loaded["S"]["A1"].value[1]
    assert block.font.b and block.font.u is None and block.font.sz is None
    loaded.close()


def test_rich_text_validation_and_xlsb(tmp_path):
    with pytest.raises(TypeError):
        CellRichText("a", 3)
    with pytest.raises(ValueError):
        InlineFont(u="wavy")
    with XlsbWriter(tmp_path / "rich.xlsb") as writer:
        writer.add_sheet("S")
        writer.write_sheet([[RICH]])
    with ExcelReader(str(tmp_path / "rich.xlsb")) as reader:
        assert list(reader.get_rows("S")) == [[str(RICH)]]


# ----------------------------------------- extra conditional formatting

EXTRA_RULES = [
    (TopBottomRule("A2:A9", rank=3, dxf=BOLD), 'type="top10" rank="3"', None),
    (TopBottomRule("A2:A9", rank=10, percent=True, bottom=True, dxf=BOLD),
     'type="top10" rank="10" percent="1" bottom="1"', None),
    (AboveAverageRule("A2:A9", below=True, std_dev=1, dxf=BOLD),
     'type="aboveAverage" aboveAverage="0" stdDev="1"', None),
    (TextRule("B2:B9", "contains", 'a"b', dxf=BOLD),
     'type="containsText" operator="containsText" text="a&quot;b"', 'NOT(ISERROR(SEARCH("a""b",B2)))'),
    (TextRule("B2:B9", "beginsWith", "x", dxf=BOLD), 'type="beginsWith" operator="beginsWith" text="x"', 'LEFT(B2,LEN("x"))="x"'),
    (DuplicateRule("B2:B9", dxf=BOLD), 'type="duplicateValues"', None),
    (DuplicateRule("B2:B9", unique=True, dxf=BOLD), 'type="uniqueValues"', None),
    (TimePeriodRule("C2:C9", "yesterday", dxf=BOLD), 'type="timePeriod" timePeriod="yesterday"',
     "FLOOR(C2,1)=TODAY()-1"),
    (BlanksRule("D2:D9", dxf=BOLD), 'type="containsBlanks"', "LEN(TRIM(D2))=0"),
    (BlanksRule("D2:D9", "no_errors", dxf=BOLD), 'type="notContainsErrors"', "NOT(ISERROR(D2))"),
]


@pytest.mark.parametrize("rule,attrs,formula", EXTRA_RULES)
def test_extra_conditional_formatting(tmp_path, rule, attrs, formula):
    path = tmp_path / "cf.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("S")
        writer.write_sheet([["a", "b", "c", "d"], [1, "x", None, None]], conditional_formats=[rule])
    xml = _sheet_xml(path)
    assert f"<cfRule {attrs} priority=" in xml and 'dxfId="0"' in xml
    if formula:
        assert f"<formula>{formula}</formula>" in xml.replace("&lt;", "<").replace("&gt;", ">")
    assert len(list(openpyxl.load_workbook(path)["S"].conditional_formatting)) == 1


def test_extra_rules_validation():
    with pytest.raises(ValueError, match="rank"):
        TopBottomRule("A1:A5", rank=200, percent=True)
    with pytest.raises(ValueError, match="operator"):
        TextRule("A1:A5", "matches", "x")
    with pytest.raises(ValueError, match="period"):
        TimePeriodRule("A1:A5", "lastYear")
    with pytest.raises(ValueError, match="kind"):
        BlanksRule("A1:A5", "empty")


# ------------------------------------------------------------ side fixes


def test_model_first_row_is_not_bold(tmp_path):
    # The writers bold first-row text by default; the model, like openpyxl, does not.
    workbook, _ = _report_workbook()
    workbook.save(tmp_path / "plain.xlsx")
    sheet = openpyxl.load_workbook(tmp_path / "plain.xlsx")["Report"]
    assert not sheet["A1"].font.b and not sheet["B1"].font.b


@pytest.mark.parametrize("writer_cls,ext", [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")])
def test_table_header_cells_may_carry_a_style(tmp_path, writer_cls, ext):
    from xlspy import Style

    rows = [[("Region", Style(bold=True)), ("Sales", "@")], ["North", 10]]
    with writer_cls(tmp_path / f"t.{ext}") as writer:
        writer.add_sheet("S")
        writer.write_sheet(rows, tables=[ExcelTable(display_name="T", ref="A1:B2")])


@pytest.mark.parametrize("writer_cls, ext", [(XlsxWriter, "xlsx"), (XlsbWriter, "xlsb")])
def test_styled_dates_keep_a_date_format(tmp_path, writer_cls, ext):
    """A date in a styled cell (range style or (value, Style)) without its
    own number format gets the default date format, not General."""
    day, moment = datetime.date(2024, 1, 2), datetime.datetime(2024, 1, 2, 3, 4)
    bold = Style(bold=True)
    path = tmp_path / f"dates.{ext}"
    with writer_cls(path) as writer:
        writer.add_sheet("S")
        rows = [["day", "moment", "ranged"]]
        # Several rows: the C encoder takes over once the xfs are recorded.
        rows += [[(day, bold), (moment, bold), day] for _ in range(5)]
        rows += [[(day, Style(bold=True, number_format="0.00")), (day, "dd/mm/yyyy"), moment]]
        writer.write_sheet(rows, styles={"C2:C7": Style(italic=True)})
    with ExcelReader(str(path)) as reader:
        values = reader.read_all("S")
    for row in values[1:6]:
        assert row == [datetime.datetime(2024, 1, 2), moment, datetime.datetime(2024, 1, 2)]
    # An explicit number format still wins.
    assert values[6][0] == 45293
    assert values[6][1:] == [datetime.datetime(2024, 1, 2), moment]
    if ext == "xlsx":
        sheet = openpyxl.load_workbook(path)["S"]
        assert [cell.number_format for cell in sheet[2]] == ["yyyy-mm-dd", "yyyy-mm-dd hh:mm", "yyyy-mm-dd"]
        assert sheet["A2"].font.b and sheet["C2"].font.i


def test_model_dates_with_a_font_keep_a_date_format(tmp_path):
    workbook = Workbook()
    ws = workbook.active
    ws["A1"] = "head"
    ws["A2"] = datetime.date(2024, 1, 2)
    ws["A2"].font = Font(bold=True)
    ws["A3"] = datetime.datetime(2024, 1, 2, 3, 4)
    assert ws["A2"].number_format == "yyyy-mm-dd"
    workbook.save(tmp_path / "new.xlsx")
    loaded = load_workbook(tmp_path / "new.xlsx")
    loaded.active["B2"] = datetime.date(2024, 5, 6)
    loaded.active["B2"].font = Font(italic=True)
    loaded.save(tmp_path / "loaded.xlsx")
    loaded.close()
    for name in ("new.xlsx", "loaded.xlsx"):
        sheet = openpyxl.load_workbook(tmp_path / name).active
        assert sheet["A2"].number_format == "yyyy-mm-dd" and sheet["A2"].font.b
        assert sheet["A3"].number_format == "yyyy-mm-dd hh:mm"
    sheet = openpyxl.load_workbook(tmp_path / "loaded.xlsx").active
    assert sheet["B2"].number_format == "yyyy-mm-dd" and sheet["B2"].font.i
