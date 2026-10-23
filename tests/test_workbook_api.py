"""Comparative tests for the openpyxl-like workbook model.

openpyxl is the independent reference implementation here:

* workbooks created by ``xlspy.Workbook`` must open in openpyxl with the same
  values, styles and sheet metadata;
* openpyxl-authored files must load into xlspy with coordinates, styles,
  merges, dimensions and metadata;
* edits saved by xlspy must be visible to openpyxl, while untouched package
  parts stay byte-identical.
"""

import datetime as dt
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.comments import Comment as OpenpyxlComment
from openpyxl.styles import (
    Alignment as OpenpyxlAlignment,
    Border as OpenpyxlBorder,
    Font as OpenpyxlFont,
    PatternFill as OpenpyxlPatternFill,
    Side as OpenpyxlSide,
)

from xlspy import (
    Alignment,
    Border,
    Comment,
    Font,
    Formula,
    Hyperlink,
    PatternFill,
    Side,
    Style,
    Workbook,
    load_workbook,
)


def _rgb(color):
    """Last six hex digits of an RGB color, or None when it has no RGB value."""
    rgb = getattr(color, "rgb", None)
    if not rgb:
        return None
    return str(rgb).upper()[-6:]


def _normalize(value):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return float(value)
    return value


def _assert_grid(actual, expected, context):
    assert set(actual) == set(expected), f"{context}: sheets differ"
    for name in expected:
        assert len(actual[name]) == len(expected[name]), f"{context} [{name}]: row count"
        for row_index, (actual_row, expected_row) in enumerate(zip(actual[name], expected[name])):
            assert len(actual_row) == len(expected_row), (
                f"{context} [{name}] row {row_index}: width differs"
            )
            for column_index, (actual_cell, expected_cell) in enumerate(zip(actual_row, expected_row)):
                assert _normalize(actual_cell) == _normalize(expected_cell), (
                    f"{context} [{name}] r{row_index}c{column_index}: "
                    f"{actual_cell!r} != {expected_cell!r}"
                )


def _xlspy_grid(workbook):
    return {
        sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)]
        for sheet in workbook.worksheets
    }


def _openpyxl_grid(path, *, data_only=True):
    book = openpyxl.load_workbook(path, data_only=data_only)
    try:
        return {
            name: [list(row) for row in book[name].iter_rows(values_only=True)]
            for name in book.sheetnames
        }
    finally:
        book.close()


def _make_openpyxl_source(path):
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A1"].font = OpenpyxlFont(bold=True, color="FFFF0000")
    ws["A1"].fill = OpenpyxlPatternFill("solid", fgColor="FFFFFF00")
    ws["A1"].alignment = OpenpyxlAlignment(horizontal="center", wrap_text=True)
    ws["A1"].border = OpenpyxlBorder(left=OpenpyxlSide(style="thin", color="FF000000"))
    ws["A1"].number_format = "#,##0"
    ws["A2"] = 1
    ws["A3"] = None
    ws["A4"] = 3
    ws["B2"] = dt.datetime(2024, 1, 2, 3, 4, 5)
    ws["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    ws["C2"] = dt.time(12, 30)
    ws["C2"].number_format = "h:mm"
    ws["D2"] = dt.timedelta(hours=27, minutes=5)
    ws["D2"].number_format = "[h]:mm"
    ws["E2"] = "=SUM(A2:A4)"
    ws["F2"] = "merge"
    ws.merge_cells("F2:G3")
    ws["H2"] = "link"
    ws["H2"].hyperlink = "https://example.com"
    ws["I2"] = "note"
    ws["I2"].comment = OpenpyxlComment("hello", "Ann")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:I4"
    ws.column_dimensions["A"].width = 22
    ws.row_dimensions[2].height = 31
    ws.sheet_properties.tabColor = "FF0000FF"
    book.create_sheet("Second")
    book["Second"]["A1"] = "untouched"
    book.properties.creator = "Comparer"
    book.properties.title = "Compare"
    book.save(path)
    book.close()
    return path


def test_new_workbook_values_formulas_and_dates(tmp_path):
    path = tmp_path / "new_values.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws["A1"] = "Item"
    ws["A2"] = 1200
    ws["A3"] = None
    ws["A4"] = 3
    ws["B1"] = "Total"
    ws["B2"] = Formula("SUM(A2:A4)", 1203)
    ws["B3"] = dt.datetime(2024, 5, 6, 7, 8, 9)
    ws["B4"] = dt.date(2024, 5, 7)
    ws["C2"] = dt.time(12, 30)
    ws["C2"].number_format = "h:mm"
    ws["D2"] = dt.timedelta(hours=27, minutes=5)
    ws["D2"].number_format = "[h]:mm"
    workbook.save(path)
    workbook.close()

    book = openpyxl.load_workbook(path, data_only=False)
    try:
        ws = book["Data"]
        assert ws["A2"].value == 1200
        assert ws["A3"].value is None
        assert ws["A4"].value == 3
        assert ws["B2"].value == "=SUM(A2:A4)"
        assert ws["B3"].value == dt.datetime(2024, 5, 6, 7, 8, 9)
        assert ws["C2"].value == dt.time(12, 30)
        assert ws["C2"].number_format == "h:mm"
        assert ws["D2"].value == dt.timedelta(hours=27, minutes=5)
        assert ws["D2"].number_format == "[h]:mm"
    finally:
        book.close()

    cached = openpyxl.load_workbook(path, data_only=True)
    try:
        assert cached["Data"]["B2"].value == 1203
    finally:
        cached.close()


def test_new_workbook_styles(tmp_path):
    path = tmp_path / "new_styles.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", start_color="FFFF00")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws["A1"].border = Border(left=Side(style="thin", color="000000"))
    ws["A1"].number_format = "#,##0"
    workbook.save(path)
    workbook.close()

    book = openpyxl.load_workbook(path)
    try:
        cell = book["Data"]["A1"]
        assert cell.font.b is True
        assert _rgb(cell.font.color) == "FF0000"
        assert cell.fill.patternType == "solid"
        assert _rgb(cell.fill.fgColor) == "FFFF00"
        assert cell.alignment.horizontal == "center"
        assert cell.alignment.vertical == "center"
        assert cell.alignment.wrap_text is True
        assert cell.border.left.style == "thin"
        assert cell.number_format == "#,##0"
    finally:
        book.close()


def test_new_workbook_structure_metadata_comments_and_links(tmp_path):
    path = tmp_path / "new_structure.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Data"
    ws["A1"] = "top"
    ws["B2"] = "merged"
    ws.merge_cells("B2:C3")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:D5"
    ws.column_dimensions["A"].width = 22
    ws.row_dimensions[2].height = 31
    ws["D1"] = Hyperlink("https://example.com", text="link")
    ws["E1"].comment = Comment("note", "Ann")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "h"
    workbook.save(path)
    workbook.close()

    book = openpyxl.load_workbook(path)
    try:
        ws = book["Data"]
        assert ws["B2"].value == "merged"
        assert sorted(str(item) for item in ws.merged_cells.ranges) == ["B2:C3"]
        assert ws.freeze_panes == "B2"
        assert ws.auto_filter.ref == "A1:D5"
        assert ws.column_dimensions["A"].width == 22
        assert ws.row_dimensions[2].height == 31
        assert ws["D1"].value == "link"
        assert ws["D1"].hyperlink.target == "https://example.com"
        assert ws["E1"].comment.text == "note"
        assert ws["E1"].comment.author == "Ann"
        assert book["Hidden"].sheet_state == "hidden"
        assert book.sheetnames == ["Data", "Hidden"]
    finally:
        book.close()


def test_openpyxl_source_loads_with_coordinates_styles_and_metadata(tmp_path):
    path = _make_openpyxl_source(tmp_path / "source.xlsx")
    workbook = load_workbook(path)
    try:
        ws = workbook["Data"]
        assert workbook.sheetnames == ["Data", "Second"]
        assert ws.dimensions == "A1:I4"
        assert ws.max_row == 4
        assert ws.max_column == 9
        assert [cell.value for cell in ws[2]] == [
            1, dt.datetime(2024, 1, 2, 3, 4, 5), dt.time(12, 30),
            dt.timedelta(hours=27, minutes=5), "=SUM(A2:A4)",
            "merge", None, "link", "note",
        ]
        assert ws["A2"].value == 1
        assert ws["A3"].value is None
        assert ws["E2"].value == "=SUM(A2:A4)"
        assert ws["E2"].data_type == "f"
        assert ws["F2"].value == "merge"
        assert ws["G3"].value is None
        assert list(ws.merged_cells.ranges) == ["F2:G3"]
        assert ws.freeze_panes == "B2"
        assert ws.auto_filter.ref == "A1:I4"
        assert ws.column_dimensions["A"].width == 22
        assert ws.row_dimensions[2].height == 31
        assert ws["A1"].font.bold is True
        assert _rgb(ws["A1"].font.color) == "FF0000"
        assert ws["A1"].fill.fill_type == "solid"
        assert _rgb(ws["A1"].fill.start_color) == "FFFF00"
        assert ws["A1"].alignment.horizontal == "center"
        assert ws["A1"].border.left.style == "thin"
        assert ws["A1"].number_format == "#,##0"
        assert ws["B2"].value == dt.datetime(2024, 1, 2, 3, 4, 5)
        assert ws["C2"].value == dt.time(12, 30)
        assert ws["D2"].value == dt.timedelta(hours=27, minutes=5)
        assert ws["H2"].hyperlink.target == "https://example.com"
        assert ws["I2"].comment.text == "hello"
        assert ws["I2"].comment.author == "Ann"
        assert _rgb(ws.sheet_properties.tabColor) == "0000FF"
        assert workbook.properties.creator == "Comparer"
        assert workbook.properties.title == "Compare"
        assert workbook["Second"]["A1"].value == "untouched"
    finally:
        workbook.close()


def test_openpyxl_blank_rows_and_chartsheets(tmp_path):
    path = tmp_path / "blank.xlsx"
    book = openpyxl.Workbook()
    ws = book.active
    ws["A1"] = 1
    ws["A3"] = 3
    book.create_chartsheet("Chart")
    book.save(path)
    book.close()

    workbook = load_workbook(path)
    try:
        assert workbook.sheetnames == ["Sheet"]
        ws = workbook["Sheet"]
        assert ws.dimensions == "A1:A3"
        assert ws.max_row == 3
        assert [cell.value for cell in ws[1]] == [1]
        assert [cell.value for cell in ws[2]] == [None]
        assert [cell.value for cell in ws[3]] == [3]
    finally:
        workbook.close()


def test_loaded_workbook_edits_are_visible_to_openpyxl(tmp_path):
    source = _make_openpyxl_source(tmp_path / "source.xlsx")
    edited = tmp_path / "edited.xlsx"
    workbook = load_workbook(source)
    try:
        ws = workbook["Data"]
        ws["A4"] = 33
        ws["A2"].font = Font(italic=True)
        ws["A1"].fill = PatternFill("solid", start_color="00FF00")
        ws.merge_cells("A6:B7")
        ws["A6"] = "new"
        ws.column_dimensions["B"].width = 35
        ws.row_dimensions[4].height = 44
        ws.freeze_panes = "C3"
        ws.auto_filter.ref = "A1:I7"
        workbook.save(edited)
    finally:
        workbook.close()

    book = openpyxl.load_workbook(edited)
    try:
        ws = book["Data"]
        assert ws["A4"].value == 33
        assert ws["A2"].font.i is True
        assert _rgb(ws["A1"].fill.fgColor) == "00FF00"
        assert sorted(str(item) for item in ws.merged_cells.ranges) == ["A6:B7", "F2:G3"]
        assert ws.freeze_panes == "C3"
        assert ws.auto_filter.ref == "A1:I7"
        assert ws.column_dimensions["B"].width == 35
        assert ws.row_dimensions[4].height == 44
        assert ws["I2"].comment.text == "hello"
        assert ws["H2"].hyperlink.target == "https://example.com"
        assert book["Second"]["A1"].value == "untouched"
    finally:
        book.close()


def test_loaded_workbook_sheet_management(tmp_path):
    source = _make_openpyxl_source(tmp_path / "source.xlsx")
    edited = tmp_path / "managed.xlsx"
    workbook = load_workbook(source)
    try:
        ws = workbook["Data"]
        ws.title = "Renamed"
        added = workbook.create_sheet("Added")
        added["A1"] = "hello"
        workbook.remove(workbook["Second"])
        assert workbook.sheetnames == ["Renamed", "Added"]
        workbook.save(edited)
    finally:
        workbook.close()

    book = openpyxl.load_workbook(edited)
    try:
        assert book.sheetnames == ["Renamed", "Added"]
        assert book["Renamed"]["A1"].value == "Header"
        assert book["Added"]["A1"].value == "hello"
        assert book.active.title == "Renamed"
    finally:
        book.close()

    with zipfile.ZipFile(edited) as archive:
        names = set(archive.namelist())
    sheets = sorted(
        name for name in names
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
    )
    assert len(sheets) == 2


def test_loaded_workbook_preserves_untouched_parts(tmp_path):
    source = _make_openpyxl_source(tmp_path / "source.xlsx")
    with zipfile.ZipFile(source, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("custom/keep.xml", "<keep>yes</keep>")
    with zipfile.ZipFile(source) as archive:
        before = {
            "custom": archive.read("custom/keep.xml"),
            "second": archive.read("xl/worksheets/sheet2.xml"),
        }

    edited = tmp_path / "preserved.xlsx"
    workbook = load_workbook(source)
    try:
        workbook["Data"]["A4"] = 33
        workbook.save(edited)
    finally:
        workbook.close()

    with zipfile.ZipFile(edited) as archive:
        assert archive.testzip() is None
        assert archive.read("custom/keep.xml") == before["custom"]
        assert archive.read("xl/worksheets/sheet2.xml") == before["second"]
    assert _openpyxl_grid(edited)["Data"][3][0] == 33


def test_read_only_streams_values_like_openpyxl(tmp_path):
    path = tmp_path / "readonly.xlsx"
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Data"
    ws.append(["name", "count", "when", "ok"])
    ws.append(["a", 1, dt.datetime(2024, 1, 2, 3, 4, 5), True])
    ws.append(["b", 2, dt.datetime(2024, 1, 3, 0, 0, 0), False])
    book.save(path)
    book.close()

    expected = _openpyxl_grid(path, data_only=True)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert _xlspy_grid(workbook) == expected
        sheet = workbook["Data"]
        assert sheet.dimensions == "A1:D3"
        cells = list(sheet.iter_rows(values_only=False))
        assert [cell.value for cell in cells[1]] == expected["Data"][1]
        assert cells[1][1].coordinate == "B2"
        assert sheet.cell(2, 1).value == "a" and sheet["B3"].value == 2
        with pytest.raises(AttributeError, match="read-only"):
            sheet["A1"] = 1
        with pytest.raises(ValueError, match="read-only"):
            workbook.save(tmp_path / "nope.xlsx")
    finally:
        workbook.close()


def test_write_only_workbook_is_readable_by_openpyxl(tmp_path):
    path = tmp_path / "write_only.xlsx"
    workbook = Workbook(write_only=True)
    ws = workbook.create_sheet("Data")
    with pytest.raises(ValueError, match="write-only"):
        ws.cell(1, 1)
    ws.append(["a", 1])
    ws.append(["b", 2])
    ws.column_dimensions["A"].width = 30
    workbook.save(path)
    workbook.close()

    assert _openpyxl_grid(path) == {"Data": [["a", 1], ["b", 2]]}
    book = openpyxl.load_workbook(path)
    try:
        assert book["Data"].column_dimensions["A"].width == 30
    finally:
        book.close()


def test_unsupported_workbook_model_paths(tmp_path):
    with pytest.raises(NotImplementedError, match="XLSB"):
        load_workbook(tmp_path / "model.xlsb")

    workbook = Workbook()
    with pytest.raises(NotImplementedError, match="named styles"):
        workbook.active["A1"].style = "Title"


def test_fast_path_modules_do_not_import_workbook_model():
    root = Path(__file__).resolve().parent.parent / "src"
    code = (
        "from xlspy import ExcelReader, XlsbWriter, XlsxWriter; "
        "import sys; "
        "found = [name for name in sys.modules "
        "if name == 'xlspy.workbook' or name.startswith('xlspy.workbook.')]; "
        "assert not found, found"
    )
    environment = dict(os.environ)
    environment["PYTHONPATH"] = str(root) + os.pathsep + environment.get("PYTHONPATH", "")
    subprocess.run(
        [sys.executable, "-c", code],
        env=environment,
        capture_output=True,
        text=True,
        check=True,
    )


def test_removing_loaded_sheet_drops_relationship_and_remaps_local_names(tmp_path):
    from openpyxl.workbook.defined_name import DefinedName as OpenpyxlDefinedName
    source = tmp_path / "source.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "A"
    book.create_sheet("B")["A1"] = 2
    third = book.create_sheet("C")
    third["A1"] = 3
    third.print_area = "A1:B2"
    third.defined_names["LocalC"] = OpenpyxlDefinedName("LocalC", attr_text="C!$A$1")
    book.save(source)
    workbook = load_workbook(source)
    workbook.remove(workbook["A"])
    target = tmp_path / "target.xlsx"
    workbook.save(target)
    workbook.close()
    with zipfile.ZipFile(target) as archive:
        rels = archive.read("xl/_rels/workbook.xml.rels")
        names = set(archive.namelist())
    # Every workbook relationship must point at an existing part.
    for target_path in re.findall(rb'Target="([^"]+)"', rels):
        part = target_path.decode().lstrip("/")
        assert (part if part.startswith("xl/") else "xl/" + part) in names
    result = openpyxl.load_workbook(target)
    assert result.sheetnames == ["B", "C"]
    assert list(result["C"].defined_names) == ["LocalC"]
    assert result["C"].print_area == "'C'!$A$1:$B$2"


def test_style_edit_keeps_values_of_cells_without_references(tmp_path):
    source, stripped = tmp_path / "source.xlsx", tmp_path / "stripped.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "S"
    book.active.append([1, 2, 3])
    book.save(source)
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(stripped, "w") as output:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<c r="[A-Z]+\d+"', b"<c", data)
            output.writestr(name, data)
    workbook = load_workbook(stripped)
    workbook["S"]["B1"].font = Font(bold=True)
    target = tmp_path / "target.xlsx"
    workbook.save(target)
    workbook.close()
    sheet = openpyxl.load_workbook(target)["S"]
    assert [cell.value for cell in sheet[1]] == [1, 2, 3]
    assert sheet["B1"].font.b



def test_read_only_sparse_rows_match_openpyxl(tmp_path):
    path = tmp_path / "sparse.xlsx"
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Data"
    ws["B2"] = "first"
    ws["D5"] = 5
    ws["A9"] = "last"
    book.save(path)
    book.close()
    reference = openpyxl.load_workbook(path, read_only=True)
    workbook = load_workbook(path, read_only=True)
    try:
        expected_sheet, sheet = reference["Data"], workbook["Data"]
        assert (sheet.min_row, sheet.min_column, sheet.max_row, sheet.max_column) == (
            expected_sheet.min_row, expected_sheet.min_column,
            expected_sheet.max_row, expected_sheet.max_column)
        for options in ({}, {"min_row": 4, "max_row": 6}, {"min_col": 2, "max_col": 4},
                        {"min_row": 3, "max_row": 12, "min_col": 1, "max_col": 2}):
            assert (list(sheet.iter_rows(values_only=True, **options))
                    == list(expected_sheet.iter_rows(values_only=True, **options)))
            assert ([[cell.value for cell in row] for row in sheet.iter_rows(**options)]
                    == [[cell.value for cell in row] for row in expected_sheet.iter_rows(**options)])
        assert sheet["D5"].value == expected_sheet["D5"].value
        sheet.reset_dimensions()
        expected_sheet.reset_dimensions()
        assert sheet.max_row is None
        assert list(sheet.values) == [tuple(row) for row in expected_sheet.values]
        assert sheet.calculate_dimension(force=True) == expected_sheet.calculate_dimension(force=True)
    finally:
        workbook.close()
        reference.close()


def test_read_only_binary_workbook_uses_worksheet_coordinates(tmp_path):
    from xlspy import XlsbUpdater, XlsbWriter
    source, path = tmp_path / "source.xlsb", tmp_path / "sparse.xlsb"
    with XlsbWriter(source) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([["a", 1]])
    updater = XlsbUpdater(source)
    try:
        updater.update_cells("Data", {"C4": 7})
        updater.save(path)
    finally:
        updater.close()
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook["Data"]
        assert sheet.dimensions == "A1:C4"
        assert list(sheet.values) == [("a", 1, None), (None,) * 3, (None,) * 3, (None, None, 7)]
        assert sheet["C4"].value == 7
    finally:
        workbook.close()


def test_serialize_like_keeps_prefixes_named_by_ignorable():
    from xml.etree import ElementTree as ET
    from xlspy._metadata import serialize_like
    original = (
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
        'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main" '
        'mc:Ignorable="x15"><sheets/><extLst><ext xmlns:xcalcf="urn:calc" uri="{1}">'
        '<xcalcf:calcFeatures/></ext></extLst></workbook>'
    )
    output = serialize_like(ET.fromstring(original), original).decode()
    assert 'xmlns:x15="http://schemas.microsoft.com/office/spreadsheetml/2010/11/main"' in output
    assert 'mc:Ignorable="x15"' in output and "<xcalcf:calcFeatures" in output
    assert "ns0" not in output and "ns1" not in output
    assert output.splitlines()[1].startswith('<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"')


def test_read_only_binary_files_written_by_xlspy_are_not_truncated(tmp_path):
    from xlspy import XlsbUpdater, XlsbWriter
    path = tmp_path / "written.xlsb"
    with XlsbWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([[1, 2, Formula("A1*2")], [4]])
    workbook = load_workbook(path, read_only=True)
    try:
        # XlsbWriter stores an A1:A1 dimension placeholder: treat it as unsized.
        assert workbook["Data"].max_row is None
        assert list(workbook["Data"].values) == [(1, 2, "=A1*2"), (4,)]
    finally:
        workbook.close()
    updated = tmp_path / "updated.xlsb"
    updater = XlsbUpdater(path)
    try:
        updater.update_cells("Data", {"B4": 9})
        updater.save(updated)
    finally:
        updater.close()
    workbook = load_workbook(updated, read_only=True, data_only=True)
    try:
        sheet = workbook["Data"]
        assert sheet.dimensions == "A1:C4"
        assert list(sheet.values) == [(1, 2, 0), (4, None, None), (None, None, None), (None, 9, None)]
    finally:
        workbook.close()
