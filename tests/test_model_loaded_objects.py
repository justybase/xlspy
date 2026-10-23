"""Objects added to the sheets of a loaded workbook (charts, images, tables,
data validation, conditional formats and print setup)."""

import re
import zipfile

import openpyxl
import pytest

from xlspy import (
    CellIsRule, Chart, ChartSeries, DataValidation, DifferentialStyle, ExcelTable, Font, Image,
    Hyperlink, PageMargins, PrintSettings, XlsxWriter, load_workbook,
)

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de"
    "0000000c4944415408d763f8cfc000000301010018dd8db40000000049454e44ae426082"
)
ROWS = [["Region", "Sales"], ["North", 10], ["South", 30], ["East", 20]]
ITALIC = DifferentialStyle(font=Font(italic=True))
BOLD = DifferentialStyle(font=Font(bold=True), number_format="0.0%")


@pytest.fixture
def load():
    """load_workbook that closes the workbooks after the test."""
    books = []

    def _load(path):
        books.append(load_workbook(path))
        return books[-1]

    yield _load
    for book in books:
        book.close()


def _source(path, *, objects=True):
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(
            ROWS,
            charts=[Chart("line", [ChartSeries("B2:B4")], anchor="H2")] if objects else None,
            tables=[ExcelTable("OldTable", "A1:B4")] if objects else None,
            conditional_formats=[CellIsRule("B2:B4", "lessThan", ["15"], dxf=ITALIC)] if objects else None,
        )
        writer.add_sheet("Other")
        writer.write_sheet([["x"]])
    return path


def _add_everything(sheet):
    sheet.add_chart(Chart("bar", [ChartSeries("B2:B4", categories="A2:A4")], title="New"), "D10")
    sheet.add_image(Image(PNG), "D30")
    sheet["D1"] = "Column"
    sheet.add_table(ExcelTable(display_name="NewTable", ref="D1:D4"))
    sheet.add_data_validation(DataValidation(type="whole", sqref="B2:B4", formula1="0", formula2="100"))
    sheet.conditional_formatting.add("B2:B4", CellIsRule("A1", "greaterThan", ["15"], dxf=BOLD))
    sheet.print_area = "A1:D4"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.oddHeader.center.text = "Report"


def _parts(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _check_relationships(parts):
    """Every relationship target exists and every new part has a content type."""
    types = parts["[Content_Types].xml"].decode()
    for name, data in parts.items():
        if not name.endswith(".rels"):
            continue
        base = name.replace("_rels/", "").removesuffix(".rels")
        directory = base.rsplit("/", 1)[0] if "/" in base else ""
        for target in re.findall(r'Target="([^"]+)"', data.decode()):
            if target.startswith("http"):
                continue
            resolved = target.lstrip("/") if target.startswith("/") else "/".join(
                part for part in (directory + "/" + target).split("/") if part
            )
            while "/../" in "/" + resolved:
                resolved = re.sub(r"[^/]+/\.\./", "", resolved, count=1)
            assert resolved in parts, (name, target)
    for name in parts:
        if name.startswith(("xl/charts/", "xl/drawings/drawing", "xl/tables/")) and not name.endswith(".rels"):
            assert f'PartName="/{name}"' in types, name


def test_loaded_workbook_saves_added_objects(tmp_path, load):
    workbook = load(_source(tmp_path / "source.xlsx"))
    _add_everything(workbook["Data"])
    target = tmp_path / "out.xlsx"
    workbook.save(target)

    parts = _parts(target)
    _check_relationships(parts)
    assert {"xl/charts/chart1.xml", "xl/charts/chart2.xml", "xl/media/image1.png",
            "xl/tables/table1.xml", "xl/tables/table2.xml"} <= set(parts)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    # The file's own rule keeps priority 1 and dxf 0; the new one follows.
    assert re.findall(r'priority="(\d+)" dxfId="(\d+)"', sheet) == [("1", "0"), ("2", "1")]
    assert '<tableParts count="2">' in sheet
    order = [sheet.index(tag) for tag in ("<conditionalFormatting", "<dataValidations", "<pageMargins",
                                          "<pageSetup", "<headerFooter", "<drawing", "<tableParts")]
    assert order == sorted(order)
    # The new table gets a free id and the drawing keeps one shape id each.
    assert re.search(r'<table\b[^>]*\bid="2"', parts["xl/tables/table2.xml"].decode())
    drawing = parts["xl/drawings/drawing1.xml"].decode()
    assert sorted(re.findall(r'cNvPr\b[^>]*?\bid="(\d+)"', drawing)) == ["1", "2", "3"]
    styles = parts["xl/styles.xml"].decode()
    assert '<dxfs count="2">' in styles
    names = parts["xl/workbook.xml"].decode()
    assert "_xlnm.Print_Area\" localSheetId=\"0\">'Data'!$A$1:$D$4<" in names
    assert "_xlnm.Print_Titles\" localSheetId=\"0\">'Data'!$1:$1<" in names

    reread = openpyxl.load_workbook(target)
    ws = reread["Data"]
    assert sorted(ws.tables) == ["NewTable", "OldTable"]
    # openpyxl reads pictures only with Pillow, which CI does not install.
    assert len(ws._charts) == 2 and drawing.count("<pic>") == 1
    assert [dv.sqref for dv in ws.data_validations.dataValidation] == ["B2:B4"] or \
        [str(dv.sqref) for dv in ws.data_validations.dataValidation] == ["B2:B4"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.oddHeader.center.text == "Report"


def test_saving_again_does_not_repeat_the_objects(tmp_path, load):
    workbook = load(_source(tmp_path / "source.xlsx"))
    _add_everything(workbook["Data"])
    first, second = tmp_path / "first.xlsx", tmp_path / "second.xlsx"
    workbook.save(first)
    workbook.save(second)
    assert sorted(_parts(first)) == sorted(_parts(second))
    assert _parts(first)["xl/worksheets/sheet1.xml"] == _parts(second)["xl/worksheets/sheet1.xml"]
    # Objects added after a save are saved the next time.
    workbook["Data"].add_chart(Chart("pie", [ChartSeries("B2:B4")]), "M2")
    workbook.save(tmp_path / "third.xlsx")
    assert "xl/charts/chart3.xml" in _parts(tmp_path / "third.xlsx")


def test_sheet_without_objects_gets_new_parts(tmp_path, load):
    workbook = load(_source(tmp_path / "plain.xlsx", objects=False))
    _add_everything(workbook["Other"])
    workbook.save(tmp_path / "out.xlsx")
    parts = _parts(tmp_path / "out.xlsx")
    _check_relationships(parts)
    sheet = parts["xl/worksheets/sheet2.xml"].decode()
    assert "<drawing " in sheet and '<tableParts xmlns="' in sheet
    assert "localSheetId=\"1\">'Other'!$A$1:$D$4<" in parts["xl/workbook.xml"].decode()
    assert len(openpyxl.load_workbook(tmp_path / "out.xlsx")["Other"]._charts) == 1


def test_added_sheet_gets_objects(tmp_path, load):
    workbook = load(_source(tmp_path / "source.xlsx"))
    sheet = workbook.create_sheet("New", 0)
    for row in ROWS:
        sheet.append(row)
    sheet.add_chart(Chart("bar", [ChartSeries("B2:B4")]), "D2")
    sheet.print_area = "A1:B4"
    workbook.save(tmp_path / "out.xlsx")
    names = _parts(tmp_path / "out.xlsx")["xl/workbook.xml"].decode()
    # The new sheet is first, so its print area is local to sheet 0.
    assert "localSheetId=\"0\">'New'!$A$1:$B$4<" in names
    assert len(openpyxl.load_workbook(tmp_path / "out.xlsx")["New"]._charts) == 1


def test_prefixed_host_sheet(tmp_path, load):
    source = _source(tmp_path / "source.xlsx")
    parts = _parts(source)
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    sheet = sheet.replace(f'xmlns="{MAIN}"', f'xmlns:x="{MAIN}"')
    sheet = re.sub(r"<(/?)([A-Za-z]\w*)(?=[\s/>])", r"<\1x:\2", sheet)
    parts["xl/worksheets/sheet1.xml"] = sheet.encode()
    prefixed = tmp_path / "prefixed.xlsx"
    with zipfile.ZipFile(prefixed, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    workbook = load(prefixed)
    _add_everything(workbook["Data"])
    workbook.save(tmp_path / "out.xlsx")
    result = _parts(tmp_path / "out.xlsx")["xl/worksheets/sheet1.xml"].decode()
    # Existing elements keep the prefix; inserted ones declare the namespace.
    assert '<x:tableParts count="2">' in result and f'<pageSetup xmlns="{MAIN}"' in result
    assert result.count("<x:conditionalFormatting") == 1
    assert result.count(f'<conditionalFormatting xmlns="{MAIN}"') == 1
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx")["Data"]
    assert sorted(ws.tables) == ["NewTable", "OldTable"] and len(ws._charts) == 2


def test_table_name_and_range_conflicts(tmp_path, load):
    workbook = load(_source(tmp_path / "source.xlsx"))
    workbook["Data"]["D1"] = "Column"
    workbook["Data"].add_table(ExcelTable(display_name="OldTable", ref="D1:D4"))
    with pytest.raises(ValueError, match="already exists"):
        workbook.save(tmp_path / "name.xlsx")
    assert not (tmp_path / "name.xlsx").exists()

    workbook = load(_source(tmp_path / "source.xlsx"))
    workbook["Data"]["C1"] = "Column"
    workbook["Data"].add_table(ExcelTable(display_name="Another", ref="B1:C4"))
    with pytest.raises(ValueError, match="overlaps"):
        workbook.save(tmp_path / "range.xlsx")


def test_print_setup_keeps_other_page_settings(tmp_path, load):
    source = tmp_path / "print.xlsx"
    with XlsxWriter(source) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS, print_settings=PrintSettings(paper_size="A4", header="&LLeft"))
    workbook = load(source)
    workbook["Data"].page_setup.orientation = "landscape"
    workbook["Data"].oddFooter.center.text = "Footer"
    workbook.save(tmp_path / "out.xlsx")
    sheet = _parts(tmp_path / "out.xlsx")["xl/worksheets/sheet1.xml"].decode()
    setup = re.search(r"<pageSetup\b[^>]*>", sheet)[0]
    assert 'paperSize="9"' in setup and 'orientation="landscape"' in setup
    footer = re.search(r"<headerFooter>.*</headerFooter>", sheet)[0]
    assert footer == "<headerFooter><oddHeader>&amp;LLeft</oddHeader><oddFooter>&amp;CFooter</oddFooter></headerFooter>"


def test_untouched_sheet_print_setup_is_not_rewritten(tmp_path, load):
    """A sheet's own print setup must not mark it as changed on every save."""
    source = tmp_path / "print.xlsx"
    with XlsxWriter(source) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([["a"]])
        writer.add_sheet("Other")
        writer.write_sheet(
            [["b"]],
            print_settings=PrintSettings(
                orientation="landscape",
                margins=PageMargins(left=0.4),
                header="&CReport",
                print_area="A1:A1",
            ),
        )
    with zipfile.ZipFile(source) as archive:
        before = archive.read("xl/worksheets/sheet2.xml")
        first = archive.read("xl/worksheets/sheet1.xml")

    workbook = load(source)
    workbook["Data"]["A1"] = "changed"
    workbook.save(tmp_path / "out.xlsx")

    parts = _parts(tmp_path / "out.xlsx")
    assert parts["xl/worksheets/sheet2.xml"] == before
    # The edited sheet really was saved.
    assert parts["xl/worksheets/sheet1.xml"] != first


def test_collapsed_summary_row_survives_row_metadata_save(tmp_path, load):
    """Excel marks a collapsed group's summary row with collapsed="1" alone;
    rewriting the row metadata must keep it."""
    plain = tmp_path / "plain.xlsx"
    with XlsxWriter(plain) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(ROWS)
    source = tmp_path / "collapsed.xlsx"
    with zipfile.ZipFile(plain) as original, zipfile.ZipFile(source, "w") as archive:
        for name in original.namelist():
            data = original.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<row r="2">', b'<row r="2" outlineLevel="1" hidden="1">')
                data = data.replace(b'<row r="3">', b'<row r="3" collapsed="1">')
            archive.writestr(name, data)

    workbook = load(source)
    workbook["Data"].row_dimensions[4].height = 30
    workbook.save(tmp_path / "out.xlsx")

    sheet = _parts(tmp_path / "out.xlsx")["xl/worksheets/sheet1.xml"].decode()
    row3 = re.search(r'<row r="3"[^>]*>', sheet)[0]
    assert 'collapsed="1"' in row3
    assert 'outlineLevel="1"' in re.search(r'<row r="2"[^>]*>', sheet)[0]


def test_failed_save_does_not_repeat_objects_on_retry(tmp_path, load):
    workbook = load(_source(tmp_path / "source.xlsx"))
    workbook["Data"].add_chart(Chart("bar", [ChartSeries("B2:B4")]), "D10")
    other = workbook["Other"]
    other.add_table(ExcelTable(display_name="OldTable", ref="A1:A1"))
    with pytest.raises(ValueError, match="already exists"):
        workbook.save(tmp_path / "failed.xlsx")
    other._tables.clear()
    workbook.save(tmp_path / "out.xlsx")
    charts = [name for name in _parts(tmp_path / "out.xlsx") if name.startswith("xl/charts/")]
    assert charts == ["xl/charts/chart1.xml", "xl/charts/chart2.xml"]


def _printed_source(path):
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet(
            [["a", "b"], [1, 2]],
            print_settings=PrintSettings(
                orientation="landscape",
                margins=PageMargins(left=0.4),
                print_area="A1:B2",
                repeat_rows="1:1",
                fit_to_width=1,
                fit_to_height=0,
                center_horizontally=True,
                header="&CReport",
                footer="&CPage &P",
            ),
        )
    return path


def test_cleared_print_settings_are_removed(tmp_path, load):
    """Print settings the file had and the sheet no longer has must leave the file."""
    workbook = load(_printed_source(tmp_path / "print.xlsx"))
    ws = workbook["Data"]
    ws.print_area = None
    ws.print_title_rows = None
    ws.page_setup.orientation = None
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_options.horizontalCentered = False
    ws.page_margins.left = 0.7
    ws.oddHeader.center.text = None
    ws.oddFooter.center.text = None
    workbook.save(tmp_path / "out.xlsx")

    parts = _parts(tmp_path / "out.xlsx")
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    names = parts["xl/workbook.xml"].decode()
    assert "_xlnm.Print_Area" not in names and "_xlnm.Print_Titles" not in names
    assert 'left="0.7"' in re.search(r"<pageMargins\b[^>]*>", sheet)[0]
    assert "orientation" not in sheet and "fitToWidth" not in sheet and "fitToPage" not in sheet
    assert "horizontalCentered" not in sheet
    assert "<oddHeader" not in sheet and "<oddFooter" not in sheet

    reread = load(tmp_path / "out.xlsx")["Data"]
    assert reread.print_area is None and reread.print_title_rows is None
    assert not reread.sheet_properties.pageSetUpPr.fitToPage
    check = openpyxl.load_workbook(tmp_path / "out.xlsx")["Data"]
    assert check.page_margins.left == 0.7 and check.page_setup.orientation is None


def test_partly_cleared_print_settings_keep_the_rest(tmp_path, load):
    workbook = load(_printed_source(tmp_path / "print.xlsx"))
    ws = workbook["Data"]
    # The file's print setup is loaded into the model.
    assert (ws.print_area, ws.print_title_rows) == ("A1:B2", "1:1")
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToHeight == 0
    assert ws.sheet_properties.pageSetUpPr.fitToPage and ws.print_options.horizontalCentered
    assert (ws.oddHeader.center.text, ws.oddFooter.center.text) == ("Report", "Page &P")
    ws.print_area = None
    ws.oddFooter.center.text = None
    workbook.save(tmp_path / "out.xlsx")

    parts = _parts(tmp_path / "out.xlsx")
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    names = parts["xl/workbook.xml"].decode()
    assert "_xlnm.Print_Area" not in names
    assert "_xlnm.Print_Titles\" localSheetId=\"0\">'Data'!$1:$1<" in names
    assert 'orientation="landscape"' in sheet and 'fitToPage="1"' in sheet
    assert "<oddHeader>&amp;CReport</oddHeader>" in sheet and "<oddFooter" not in sheet


def _linked_source(path):
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        rows = [[f"r{row}", None] for row in range(1, 13)]
        rows[9][0] = Hyperlink("https://example.com/a", text="A link", tooltip="Tip")
        rows[2][1] = Hyperlink("#Data!A1", text="Back")
        rows[11][1] = Hyperlink("https://example.com/b")
        writer.write_sheet(rows)
        writer.add_sheet("Other")
        writer.write_sheet([["x"]])
    return path


def _links(path, sheet="Data"):
    book = openpyxl.load_workbook(path)
    try:
        return {
            cell.coordinate: (cell.hyperlink.target or "#" + cell.hyperlink.location,
                              cell.hyperlink.display, cell.hyperlink.tooltip)
            for row in book[sheet].iter_rows() for cell in row if cell.hyperlink is not None
        }
    finally:
        book.close()


def _hyperlink_rels(parts, number=1):
    rels = parts.get(f"xl/worksheets/_rels/sheet{number}.xml.rels", b"").decode()
    return re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', rels)


def test_hyperlinks_move_with_inserted_and_deleted_rows(tmp_path, load):
    workbook = load(_linked_source(tmp_path / "links.xlsx"))
    ws = workbook["Data"]
    ws.insert_rows(2)
    workbook.save(tmp_path / "inserted.xlsx")
    assert _links(tmp_path / "inserted.xlsx") == {
        "A11": ("https://example.com/a", None, "Tip"),
        "B4": ("#Data!A1", None, None),
        "B13": ("https://example.com/b", None, None),
    }
    check = openpyxl.load_workbook(tmp_path / "inserted.xlsx")["Data"]
    assert (check["A11"].value, check["B4"].value, check["A10"].value) == ("A link", "Back", "r9")
    parts = _parts(tmp_path / "inserted.xlsx")
    _check_relationships(parts)
    # The replaced links' relationships do not linger.
    assert sorted(_hyperlink_rels(parts)) == ["https://example.com/a", "https://example.com/b"]

    ws.delete_rows(1, 3)
    ws.move_range("A8", rows=1, cols=2)
    workbook.save(tmp_path / "deleted.xlsx")
    assert _links(tmp_path / "deleted.xlsx") == {
        "C9": ("https://example.com/a", None, "Tip"),
        "B1": ("#Data!A1", None, None),
        "B10": ("https://example.com/b", None, None),
    }


def test_hyperlinks_set_and_removed_on_a_loaded_sheet(tmp_path, load):
    workbook = load(_linked_source(tmp_path / "links.xlsx"))
    ws = workbook["Data"]
    ws["A1"].hyperlink = Hyperlink("https://example.com/new", tooltip="New")
    ws["A10"].hyperlink = None
    ws["B3"].hyperlink = Hyperlink("#Other!A1")
    new = workbook.create_sheet("Added")
    new["A1"] = Hyperlink("https://example.com/added", text="Added")
    workbook.save(tmp_path / "out.xlsx")

    assert _links(tmp_path / "out.xlsx") == {
        "A1": ("https://example.com/new", None, "New"),
        "B3": ("#Other!A1", None, None),
        "B12": ("https://example.com/b", None, None),
    }
    assert _links(tmp_path / "out.xlsx", "Added") == {
        "A1": ("https://example.com/added", "Added", None),
    }
    parts = _parts(tmp_path / "out.xlsx")
    _check_relationships(parts)
    assert sorted(_hyperlink_rels(parts)) == ["https://example.com/b", "https://example.com/new"]
    sheet = parts["xl/worksheets/sheet1.xml"].decode()
    order = [sheet.index(tag) for tag in ("</sheetData>", "<hyperlinks>", "<pageMargins")]
    assert order == sorted(order)
