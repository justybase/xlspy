"""Interoperability and preservation tests for classic notes and template edits."""

import datetime
import io
import re
import zipfile
from xml.etree import ElementTree as ET

import pytest
import openpyxl
from openpyxl.comments import Comment as OpenpyxlComment
from openpyxl.styles import Font as OpenpyxlFont

from xlspy import (
    Comment, DefinedName, ExcelReader, Formula, SheetProtection, Style,
    WorkbookProtection, XlsbUpdater, XlsbWriter, XlsmUpdater, XlsxUpdater, XlsxWriter,
)
from xlspy._metadata import MAIN, REL, PKG_REL, CT, password_hash
from xlspy._text import xstring_unescape
from xlspy.comments import X, V, find_comment_parts
from xlspy.updater_utils import ZipPackage


def make_book(path, **options):
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([["Header"], [10]], **options)
        writer.add_sheet("Other")
        writer.write_sheet([["unchanged"]])


def payloads(path):
    with zipfile.ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def replace_parts(path, changes):
    parts = payloads(path)
    parts.update(changes)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def get_comment_parts(path):
    with zipfile.ZipFile(path) as archive:
        return find_comment_parts(archive.read, archive.namelist(), "xl/worksheets/sheet1.xml")[2]


def test_comments_writer_interoperability_and_dimensions(tmp_path):
    path = tmp_path / "notes.xlsx"
    comments = {"A1": Comment("Zażółć <&>\n next _x0041_", "Jan & Ola", 300, 120),
                "C8": Comment("empty cell", "Jan & Ola"), "B2": Comment("second", "Żaneta")}
    make_book(path, comments=comments)
    with ExcelReader(str(path)) as reader:
        assert reader.get_comments("Data") == {ref: Comment(c.text, c.author) for ref, c in comments.items()}
        assert reader.get_comments("Other") == {}
    book = openpyxl.load_workbook(path)
    try:
        for ref, comment in comments.items():
            # openpyxl does not decode ST_Xstring escapes in comment text.
            assert xstring_unescape(book["Data"][ref].comment.text) == comment.text
            assert book["Data"][ref].comment.author == comment.author
    finally:
        book.close()
    parts = payloads(path)
    vml = ET.fromstring(parts[get_comment_parts(path)["vmlDrawing"][0]])
    shape = next(vml.iter(f"{{{V}}}shape"))
    assert "width:225pt" in shape.attrib["style"]
    assert "height:90pt" in shape.attrib["style"]
    assert len(ET.fromstring(parts[get_comment_parts(path)["comments"][0]]).find(f"{{{MAIN}}}authors")) == 2


def test_openpyxl_comments_read_update_and_remove(tmp_path):
    path = tmp_path / "existing.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Data"
    sheet["A1"] = "header"
    sheet["A1"].comment = OpenpyxlComment("original", "Alice")
    sheet["D9"].comment = OpenpyxlComment("keep me", "Bob")
    book.save(path)
    book.close()
    updater = XlsxUpdater(path)
    try:
        assert updater.get_comments("Data")["A1"] == Comment("original", "Alice")
        updater.set_comments("Data", {"A1": Comment("changed", "Jan"), "B3": Comment("new", "Ola")})
        updater.set_comments("Data", {"B3": None})
        updater.replace_sheet_data("Data", [["replacement"]])
        updater.save()
    finally:
        updater.close()
    book = openpyxl.load_workbook(path)
    try:
        assert book["Data"]["A1"].comment.text == "changed"
        assert book["Data"]["D9"].comment.text == "keep me"
        assert book["Data"]["B3"].comment is None
    finally:
        book.close()


def test_remove_last_comment_cleans_parts_and_references(tmp_path):
    path = tmp_path / "last.xlsx"
    make_book(path, comments={"A1": Comment("last", "Jan")})
    old_parts = get_comment_parts(path)
    updater = XlsxUpdater(path)
    try:
        updater.set_comments("Data", {"A1": None})
        updater.save()
    finally:
        updater.close()
    parts = payloads(path)
    assert old_parts["comments"][0] not in parts
    assert old_parts["vmlDrawing"][0] not in parts
    assert ET.fromstring(parts["xl/worksheets/sheet1.xml"]).find(f"{{{MAIN}}}legacyDrawing") is None
    assert get_comment_parts(path) == {}
    assert old_parts["comments"][0].encode() not in parts["[Content_Types].xml"]


def test_vml_controls_and_untouched_notes_are_preserved(tmp_path):
    path = tmp_path / "controls.xlsx"
    make_book(path, comments={"A1": Comment("remove", "Jan"), "B2": Comment("keep", "Ola")})
    vp = get_comment_parts(path)["vmlDrawing"][0]
    control = (f'<v:shape xmlns:v="{V}" xmlns:x="{X}" id="control1" type="#custom">'
               '<x:ClientData ObjectType="Button"><x:Macro>run_me</x:Macro></x:ClientData></v:shape>').encode()
    original = payloads(path)[vp]
    replace_parts(path, {vp: original.replace(b"</xml>", control + b"</xml>")})
    updater = XlsxUpdater(path)
    try:
        updater.set_comments("Data", {"A1": None, "C3": Comment("new", "Jan")})
        updater.set_comments("Data", {"B2": None, "C3": None})
        updater.save()
    finally:
        updater.close()
    parts = payloads(path)
    assert control in parts[vp]
    assert get_comment_parts(path).keys() == {"vmlDrawing"}
    assert ET.fromstring(parts["xl/worksheets/sheet1.xml"]).find(f"{{{MAIN}}}legacyDrawing") is not None


def test_malformed_vml_is_atomic(tmp_path):
    path = tmp_path / "bad.xlsx"
    make_book(path, comments={"A1": Comment("old", "Jan")})
    vp = get_comment_parts(path)["vmlDrawing"][0]
    original = payloads(path)[vp]
    replace_parts(path, {vp: original.replace(b"<x:Row>0</x:Row>", b"")})
    updater = XlsxUpdater(path)
    try:
        before = updater.to_bytes()
        with pytest.raises(ValueError, match="VML"):
            updater.set_comments("Data", {"A1": Comment("new", "Jan")})
        assert updater.to_bytes() == before
    finally:
        updater.close()


def test_xlsm_vba_and_unrelated_parts_preserved(tmp_path):
    path = tmp_path / "macros.xlsm"
    make_book(path)
    replace_parts(path, {"xl/vbaProject.bin": b"opaque VBA", "customXml/item1.xml": b"<opaque/>"})
    before = payloads(path)
    updater = XlsmUpdater(path)
    try:
        updater.set_comments("Data", {"B2": Comment("note", "Jan")})
        updater.update_cells("Data", {"A2": 42})
        updater.save()
    finally:
        updater.close()
    after = payloads(path)
    for name in ("xl/vbaProject.bin", "customXml/item1.xml", "xl/worksheets/sheet2.xml", "xl/styles.xml"):
        assert after[name] == before[name]


@pytest.mark.parametrize("ref", ["A0", "A1:B2", "XFE1", "A1048577", "Data!A1"])
def test_invalid_comment_addresses(tmp_path, ref):
    with XlsxWriter(tmp_path / "invalid.xlsx") as writer:
        with pytest.raises(ValueError):
            writer.write_sheet([[1]], comments={ref: Comment("text", "author")})
        writer.write_sheet([[1]])


def test_merged_comment_anchor_and_duplicate_refs(tmp_path):
    with XlsxWriter(tmp_path / "merged.xlsx") as writer:
        with pytest.raises(ValueError, match="top-left"):
            writer.write_sheet([[1]], merges=["A1:B2"], comments={"B2": Comment("text", "author")})
        with pytest.raises(ValueError, match="duplicate"):
            writer.write_sheet([[1]], comments={"a1": Comment("a", "a"), "$A$1": Comment("b", "b")})
        writer.write_sheet([[1]], merges=["A1:B2"], comments={"A1": Comment("ok", "author")})


@pytest.mark.parametrize("kwargs", [{"text": 12}, {"author": None}, {"width": 0}, {"height": float("nan")}, {"width": True}])
def test_comment_validation(kwargs):
    args = {"text": "text", "author": "author"} | kwargs
    with pytest.raises((TypeError, ValueError)):
        Comment(**args)


def test_sparse_updates_types_style_comment_and_untouched_xml(tmp_path):
    path = tmp_path / "sparse.xlsx"
    make_book(path, comments={"A2": Comment("keep", "Jan")}, styles={"A2": Style(bold=True)})
    before = payloads(path)
    updater = XlsxUpdater(path)
    try:
        updater.update_cells("Data", {"A2": 123, "C2": " a & <b> _x0041_\r\n", "B1": True, "D4": 3.25})
        updater.save()
    finally:
        updater.close()
    after = payloads(path)
    before_root = ET.fromstring(before["xl/worksheets/sheet1.xml"])
    after_root = ET.fromstring(after["xl/worksheets/sheet1.xml"])
    assert ET.tostring(before_root.find(f"{{{MAIN}}}sheetData/{{{MAIN}}}row/{{{MAIN}}}c")) == ET.tostring(after_root.find(f"{{{MAIN}}}sheetData/{{{MAIN}}}row/{{{MAIN}}}c"))
    for name in before:
        if name != "xl/worksheets/sheet1.xml":
            assert before[name] == after[name]
    book = openpyxl.load_workbook(path)
    try:
        sheet = book["Data"]
        assert sheet["A2"].value == 123
        assert sheet["A2"].font.bold
        assert sheet["A2"].comment.text == "keep"
        assert sheet["B1"].value is True
        assert sheet["D4"].value == 3.25
    finally:
        book.close()
    with ExcelReader(str(path)) as reader:
        rows = reader.read_all("Data")
        assert rows[1][2] == " a & <b> _x0041_\r\n"


def test_clear_value_keeps_style_and_note(tmp_path):
    path = tmp_path / "clear.xlsx"
    make_book(path, comments={"A2": Comment("keep", "Jan")}, styles={"A2": Style(bold=True)})
    updater = XlsxUpdater(path)
    try:
        updater.update_cells("Data", {"A2": None})
        updater.save()
    finally:
        updater.close()
    book = openpyxl.load_workbook(path)
    try:
        assert book["Data"]["A2"].value is None
        assert book["Data"]["A2"].font.bold
        assert book["Data"]["A2"].comment.text == "keep"
    finally:
        book.close()


def test_sparse_updates_formulas_and_calc_chain(tmp_path):
    path = tmp_path / "formulas.xlsx"
    make_book(path)
    parts = payloads(path)
    rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
    ET.SubElement(rels, f"{{{PKG_REL}}}Relationship", {"Id": "chain", "Type": REL + "/calcChain", "Target": "calcChain.xml"})
    types = ET.fromstring(parts["[Content_Types].xml"])
    ET.SubElement(types, f"{{{CT}}}Override", {"PartName": "/xl/calcChain.xml", "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"})
    replace_parts(path, {"xl/_rels/workbook.xml.rels": ET.tostring(rels), "[Content_Types].xml": ET.tostring(types), "xl/calcChain.xml": b'<calcChain/>'})
    updater = XlsxUpdater(path)
    try:
        updater.update_cells("Data", {"B2": Formula("A2*2", 20), "C2": Formula('"ok"', "ok")})
        updater.save()
    finally:
        updater.close()
    parts = payloads(path)
    assert "xl/calcChain.xml" not in parts
    assert b"calcChain" not in parts["xl/_rels/workbook.xml.rels"]
    assert b"calcChain" not in parts["[Content_Types].xml"]
    assert ET.fromstring(parts["xl/workbook.xml"]).find(f"{{{MAIN}}}calcPr").get("forceFullCalc") == "1"
    with ExcelReader(str(path)) as reader:
        assert reader.read_all("Data")[1] == [10, 20, "ok"]
        assert reader.read_all("Data", data_only=False)[1] == [10, "=A2*2", '="ok"']


def test_sparse_updates_are_atomic(tmp_path):
    path = tmp_path / "atomic.xlsx"
    make_book(path)
    updater = XlsxUpdater(path)
    try:
        before = updater.to_bytes()
        for changes in ({"A2": 5, "B2": object()}, {"A2": 5, "B2": float("nan")}, {"A2": 5, "XFE2": 1}):
            with pytest.raises((ValueError, TypeError)):
                updater.update_cells("Data", changes)
            assert updater.to_bytes() == before
    finally:
        updater.close()


def test_names_and_protection_writer_and_updater(tmp_path):
    path = tmp_path / "metadata.xlsx"
    with XlsxWriter(path) as writer:
        writer.add_sheet("Data")
        writer.write_sheet([[1]], sheet_protection=SheetProtection(password="abc", select_locked_cells=False))
        writer.set_defined_name(DefinedName("Rates", "'Data'!$A$1"))
        writer.set_defined_name(DefinedName("Rates", "42", "Data"))
        writer.set_workbook_protection(WorkbookProtection(password="abc"))
    book = openpyxl.load_workbook(path)
    try:
        assert book.defined_names["Rates"].attr_text == "'Data'!$A$1"
        assert book["Data"].defined_names["Rates"].attr_text == "42"
        assert book["Data"].protection.sheet
        assert book["Data"].protection.selectLockedCells
        assert book["Data"].protection.password == password_hash("abc")
        assert book.security.lockStructure
        assert book.security.workbookPassword == password_hash("abc")
    finally:
        book.close()
    updater = XlsxUpdater(path)
    try:
        updater.set_defined_name(DefinedName("rates", "99"))
        updater.remove_defined_name("Rates", "Data")
        assert updater.get_defined_names() == [DefinedName("rates", "99")]
        updater.set_sheet_protection("Data", SheetProtection(enabled=False))
        updater.set_workbook_protection(WorkbookProtection(lock_structure=False))
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.get_defined_names() == [DefinedName("rates", "99")]
    book = openpyxl.load_workbook(path)
    try:
        assert not book["Data"].protection.sheet
        assert not book.security or not book.security.lockStructure
    finally:
        book.close()


@pytest.mark.parametrize("name", ["A1", "R1C1", "C", "R", "_xlnm.Print_Area", "has space", "1name", "a" * 256])
def test_invalid_names(name):
    with pytest.raises(ValueError):
        DefinedName(name, "0")


def test_package_remove_rollback_and_commit(tmp_path):
    path = tmp_path / "package.xlsx"
    make_book(path)
    package = ZipPackage(path)
    try:
        before = package.to_bytes()
        package.begin_transaction()
        package.remove("xl/styles.xml")
        package.set("new.xml", b"new")
        package.rollback_transaction()
        assert package.to_bytes() == before
        package.begin_transaction()
        package.remove("xl/styles.xml")
        package.commit_transaction()
        with zipfile.ZipFile(io.BytesIO(package.to_bytes())) as archive:
            assert "xl/styles.xml" not in archive.namelist()
            assert archive.read("xl/workbook.xml")
    finally:
        package.close()


def test_xlsb_metadata_features_are_available(tmp_path):
    path = tmp_path / "supported.xlsb"
    with XlsbWriter(path) as writer:
        writer.write_sheet([[1]], comments={"A1": Comment("a", "b")})
        writer.set_defined_name(DefinedName("Name", "0"))
    with ExcelReader(str(path)) as reader:
        assert reader.get_comments("Sheet1") == {"A1": Comment("a", "b")}
        assert reader.read_all("Sheet1", data_only=False) == [[1]]
    updater = XlsbUpdater(path)
    try:
        updater.update_cells("Sheet1", {"A1": 2})
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.read_all('Sheet1') == [[2]]


@pytest.mark.parametrize("streamed", [False, True])
def test_notes_survive_shortened_data_region(tmp_path, streamed):
    path = tmp_path / "short.xlsx"
    make_book(path, comments={"D20": Comment("keep outside range", "Jan")})
    updater = XlsxUpdater(path)
    try:
        method = updater.replace_sheet_data_stream if streamed else updater.replace_sheet_data
        method("Data", [[1]])
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.get_comments("Data")["D20"].text == "keep outside range"


@pytest.mark.parametrize("streamed", [False, True])
def test_sparse_edits_interleave_with_data_replacement(tmp_path, streamed):
    path = tmp_path / "mixed.xlsx"
    make_book(path)
    updater = XlsxUpdater(path)
    try:
        method = updater.replace_sheet_data_stream if streamed else updater.replace_sheet_data
        method("Data", [["first"], ["second"]])
        updater.update_cells("Data", {"A1": "inline"})
        method("Other", [["third"]])
        updater.set_comments("Data", {"C5": Comment("still here", "Jan")})
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.read_all("Data") == [["inline"], ["second"]]
        assert reader.read_all("Other") == [["third"]]
    parts = payloads(path)
    count = int(ET.fromstring(parts["xl/sharedStrings.xml"]).get("count"))
    actual = sum(len(ET.fromstring(parts[name]).findall(f'.//{{{MAIN}}}c[@t="s"]')) for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"))
    assert count == actual


def test_package_failure_rolls_back_all_comment_parts(tmp_path, monkeypatch):
    path = tmp_path / "rollback.xlsx"
    make_book(path)
    updater = XlsxUpdater(path)
    try:
        before = updater.to_bytes()
        original = updater._package.set
        def failing_set(name, data):
            if name == "[Content_Types].xml":
                raise OSError("simulated staging failure")
            return original(name, data)
        monkeypatch.setattr(updater._package, "set", failing_set)
        with pytest.raises(OSError, match="simulated"):
            updater.set_comments("Data", {"A1": Comment("new", "Jan")})
        assert updater.to_bytes() == before
    finally:
        updater.close()


@pytest.mark.parametrize("mac_epoch", [False, True])
def test_sparse_date_edits_reuse_existing_style(tmp_path, mac_epoch):
    path = tmp_path / "dates.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Data"
    if mac_epoch:
        from openpyxl.utils.datetime import MAC_EPOCH
        book.epoch = MAC_EPOCH
    book.active["A1"] = datetime.datetime(2025, 1, 1)
    book.save(path)
    book.close()
    updater = XlsxUpdater(path)
    try:
        updater.update_cells("Data", {"A1": datetime.datetime(2026, 10, 20, 12, 30)})
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.read_all("Data") == [[datetime.datetime(2026, 10, 20, 12, 30)]]


def test_sparse_empty_sheet_and_missing_dimension(tmp_path):
    path = tmp_path / "empty.xlsx"
    make_book(path)
    replace_parts(path, {"xl/worksheets/sheet1.xml": f'<worksheet xmlns="{MAIN}"><sheetData/></worksheet>'.encode()})
    updater = XlsxUpdater(path)
    try:
        updater.update_cells("Data", {"D9": 42, "B2": "new"})
        updater.update_cells("Data", {"A1": True})
        updater.save()
    finally:
        updater.close()
    book = openpyxl.load_workbook(path)
    try:
        assert book["Data"]["D9"].value == 42
        assert book["Data"]["B2"].value == "new"
        assert book["Data"]["A1"].value is True
    finally:
        book.close()


def test_dataframe_comments_forwarded(tmp_path):
    pandas = pytest.importorskip("pandas")
    from xlspy import to_xlsx
    path = tmp_path / "dataframe.xlsx"
    to_xlsx(pandas.DataFrame({"amount": [42]}), path, comments={"A2": Comment("ok", "Jan")}, sheet_protection=SheetProtection())
    with ExcelReader(str(path)) as reader:
        assert reader.get_comments("Sheet1")["A2"].text == "ok"


def test_threaded_comment_parts_preserved_and_collision_refused(tmp_path):
    path = tmp_path / "threaded.xlsx"
    make_book(path)
    part = "xl/threadedComments/threadedComment1.xml"
    threaded = b'<ThreadedComments xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments"><threadedComment ref="D4" personId="person1" id="thread1"><text>conversation</text></threadedComment></ThreadedComments>'
    rels = f'<Relationships xmlns="{PKG_REL}"><Relationship Id="threaded" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" Target="../threadedComments/threadedComment1.xml"/></Relationships>'
    replace_parts(path, {part: threaded, "xl/worksheets/_rels/sheet1.xml.rels": rels.encode()})
    updater = XlsxUpdater(path)
    try:
        before = updater.to_bytes()
        with pytest.raises(ValueError, match="threaded"):
            updater.set_comments("Data", {"D4": Comment("collision", "Jan")})
        assert updater.to_bytes() == before
        updater.set_comments("Data", {"A1": Comment("separate", "Jan")})
        updater.save()
    finally:
        updater.close()
    assert payloads(path)[part] == threaded


def test_comment_edits_preserve_combined_report_parts(tmp_path):
    from test_xlsx_integration import build_report, _png
    path = build_report(tmp_path / "report.xlsx", _png(tmp_path / "logo.png"))
    before = payloads(path)
    updater = XlsxUpdater(path)
    try:
        updater.set_comments("Summary", {"B3": Comment("check quantity", "Jan")})
        updater.update_cells("Summary", {"B3": 125})
        updater.save()
    finally:
        updater.close()
    after = payloads(path)
    changed = {"[Content_Types].xml", "xl/worksheets/sheet1.xml", "xl/worksheets/_rels/sheet1.xml.rels"}
    for name in before.keys() - changed:
        assert before[name] == after[name], name
    # The picture survives in the package; openpyxl only loads images when
    # Pillow is installed, which the test extras do not require.
    assert any(name.startswith("xl/media/") for name in after)
    assert any(re.search(rb"<(?:\w+:)?pic[\s>/]", data) for name, data in after.items() if name.startswith("xl/drawings/drawing"))
    try:
        import PIL  # noqa: F401
    except ImportError:
        has_pillow = False
    else:
        has_pillow = True
    book = openpyxl.load_workbook(path)
    try:
        assert len(book["Summary"]._charts) == 3
        if has_pillow:
            assert len(book["Summary"]._images) == 1
        assert book["Summary"]["B3"].comment.text == "check quantity"
    finally:
        book.close()


@pytest.mark.parametrize("kind", ["shared", "array", "dataTable"])
def test_sparse_group_formula_edits_refused_atomically(tmp_path, kind):
    path = tmp_path / "group.xlsx"
    make_book(path)
    follower = '<f t="shared" si="0"/>' if kind == "shared" else ""
    source = f'<worksheet xmlns="{MAIN}"><sheetData><row r="1"><c r="B1"><f t="{kind}" si="0" ref="B1:B2">A1*2</f><v>2</v></c></row><row r="2"><c r="B2">{follower}<v>4</v></c></row></sheetData></worksheet>'
    replace_parts(path, {"xl/worksheets/sheet1.xml": source.encode()})
    updater = XlsxUpdater(path)
    try:
        before = updater.to_bytes()
        with pytest.raises(ValueError, match="unsupported"):
            updater.update_cells("Data", {"B2": 9})
        assert updater.to_bytes() == before
    finally:
        updater.close()


def test_prefixed_workbook_sheet_styles_and_comments(tmp_path):
    path = tmp_path / "prefixes.xlsx"
    make_book(path, comments={"A1": Comment("original", "Jan")})
    parts = payloads(path)
    replacements = {name: ET.tostring(ET.fromstring(data), encoding="utf-8") for name, data in parts.items() if name.endswith(".xml") or name.endswith(".rels")}
    replace_parts(path, replacements)
    updater = XlsxUpdater(path)
    try:
        updater.set_comments("Data", {"A1": Comment("updated", "Jan")})
        updater.update_cells("Data", {"A2": 45, "B2": Formula("A2*2", 90)})
        updater.set_defined_name(DefinedName("Amount", "'Data'!$A$2"))
        updater.save()
    finally:
        updater.close()
    with ExcelReader(str(path)) as reader:
        assert reader.read_all("Data") == [["Header"], [45, 90]]
        assert reader.get_comments("Data")["A1"].text == "updated"
        assert reader.get_defined_names() == [DefinedName("Amount", "'Data'!$A$2")]


def test_notes_do_not_use_header_footer_vml(tmp_path):
    path = tmp_path / "hf.xlsx"
    make_book(path)
    hf_vml = (f'<xml xmlns:v="{V}"><v:shape id="LH" style="width:10pt;height:10pt"/></xml>').encode()
    sheet = payloads(path)["xl/worksheets/sheet1.xml"].replace(b"</worksheet>", b'<legacyDrawingHF r:id="rIdHF"/></worksheet>')
    replace_parts(path, {
        "xl/worksheets/sheet1.xml": sheet,
        "xl/drawings/vmlDrawingHF1.vml": hf_vml,
        "xl/worksheets/_rels/sheet1.xml.rels": (
            f'<Relationships xmlns="{PKG_REL}"><Relationship Id="rIdHF" '
            f'Type="{REL}/vmlDrawing" Target="../drawings/vmlDrawingHF1.vml"/></Relationships>'
        ).encode(),
    })
    added, removed = tmp_path / "added.xlsx", tmp_path / "removed.xlsx"
    for source, target, value in ((path, added, Comment("note", "Ann")), (added, removed, None)):
        updater = XlsxUpdater(source)
        try:
            updater.set_comments("Data", {"A2": value})
            updater.save(target)
        finally:
            updater.close()
    parts = payloads(added)
    assert parts["xl/drawings/vmlDrawingHF1.vml"] == hf_vml
    legacy = ET.fromstring(parts["xl/worksheets/sheet1.xml"]).find(f"{{{MAIN}}}legacyDrawing")
    assert legacy.get(f"{{{REL}}}id") != "rIdHF"
    with ExcelReader(str(added)) as reader:
        assert reader.get_comments("Data") == {"A2": Comment("note", "Ann")}
    parts = payloads(removed)
    assert parts["xl/drawings/vmlDrawingHF1.vml"] == hf_vml
    assert b'Id="rIdHF"' in parts["xl/worksheets/_rels/sheet1.xml.rels"]
