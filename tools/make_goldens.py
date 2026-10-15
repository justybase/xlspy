"""Regenerate the golden workbooks under ``tests/data/``.

Goldens are produced by *independent* writers (openpyxl, xlsxwriter), so xlspy
is validated against files it did not create itself.  Every ``golden_*``
workbook is accompanied by a ``.json`` sidecar with the expected xlspy reading;
``tests/test_cross_validation.py`` discovers the files automatically.

Usage:
    python tools/make_goldens.py

To add a golden produced by Microsoft Excel (for example an ``.xlsb`` file), drop
the workbook into ``tests/data/`` together with a sidecar of the shape
``{"sheets": {"<name>": [[cell, ...], ...]}}`` where cells are JSON scalars or
``{"datetime": "<ISO>"}``.  No code changes are needed.
"""

import json
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from golden_data import (  # noqa: E402
    DATETIME_FORMAT,
    DATE_FORMAT,
    SHEETS,
    SHORT_DATE_FORMAT,
    encoded_sheets,
    sheets_1904,
)

DATA_DIR = ROOT / "tests" / "data"


def _write_sidecar(path, sheets):
    sidecar = path.with_suffix(path.suffix + ".json")
    payload = {"sheets": encoded_sheets(sheets)}
    sidecar.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def _write_openpyxl(path, *, date_1904=False):
    import openpyxl
    from openpyxl.utils.datetime import CALENDAR_MAC_1904

    workbook = openpyxl.Workbook()
    if date_1904:
        workbook.epoch = CALENDAR_MAC_1904
    sheets = sheets_1904() if date_1904 else SHEETS
    for index, (name, rows) in enumerate(sheets.items()):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = name
        header = rows[0]
        for column, title in enumerate(header, 1):
            sheet.cell(row=1, column=column, value=title)
        for row_index, row in enumerate(rows[1:], 2):
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row=row_index, column=column, value=value)
                title = header[column - 1]
                if title == "date_iso":
                    cell.number_format = DATE_FORMAT
                elif title == "date_short":
                    cell.number_format = SHORT_DATE_FORMAT
                elif title == "datetime":
                    cell.number_format = DATETIME_FORMAT
    workbook.save(path)
    workbook.close()


def _write_xlsxwriter(path):
    import xlsxwriter

    workbook = xlsxwriter.Workbook(str(path))
    formats = {
        "date_iso": workbook.add_format({"num_format": DATE_FORMAT}),
        "date_short": workbook.add_format({"num_format": SHORT_DATE_FORMAT}),
        "datetime": workbook.add_format({"num_format": DATETIME_FORMAT}),
    }
    for name, rows in SHEETS.items():
        sheet = workbook.add_worksheet(name)
        header = rows[0]
        for column, value in enumerate(header):
            sheet.write(0, column, value)
        for row_index, row in enumerate(rows[1:], 1):
            for column, value in enumerate(row):
                title = header[column]
                if title in formats and value is not None:
                    sheet.write_datetime(row_index, column, value, formats[title])
                else:
                    sheet.write(row_index, column, value)
    workbook.close()


def main():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    for name, writer, sheets in (
        ("golden_openpyxl.xlsx", _write_openpyxl, SHEETS),
        (
            "golden_openpyxl_1904.xlsx",
            lambda path: _write_openpyxl(path, date_1904=True),
            sheets_1904(),
        ),
        ("golden_xlsxwriter.xlsx", _write_xlsxwriter, SHEETS),
    ):
        path = DATA_DIR / name
        writer(path)
        _write_sidecar(path, sheets)
        print(f"wrote {path.relative_to(ROOT)} and {path.name}.json")


if __name__ == "__main__":
    main()
